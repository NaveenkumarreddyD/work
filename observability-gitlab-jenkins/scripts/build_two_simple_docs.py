from pathlib import Path
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

DIRP = "/Users/naveendubba/Documents/Ibm mas and manage/"
BLUE="1F4E78"; BLUE2="2E74B5"; DARK="1F2937"; MUTED="667085"
LIGHT_BLUE="E8EEF5"; CODE_FILL="F5F7FA"; CAUTION_FILL="FFF4CE"; GREEN_FILL="EAF5EA"; RISK_FILL="FDECEC"
TW=9360; TI=120


def cshade(cell, fill):
    tp=cell._tc.get_or_add_tcPr(); shd=tp.find(qn("w:shd"))
    if shd is None: shd=OxmlElement("w:shd"); tp.append(shd)
    shd.set(qn("w:fill"), fill)

def cmargins(cell,t=80,s=120,b=80,e=120):
    tp=cell._tc.get_or_add_tcPr(); tm=tp.first_child_found_in("w:tcMar")
    if tm is None: tm=OxmlElement("w:tcMar"); tp.append(tm)
    for m,v in (("top",t),("start",s),("bottom",b),("end",e)):
        n=tm.find(qn(f"w:{m}"))
        if n is None: n=OxmlElement(f"w:{m}"); tm.append(n)
        n.set(qn("w:w"),str(v)); n.set(qn("w:type"),"dxa")

def geom(table,widths):
    assert sum(widths)==TW
    table.autofit=False; table.alignment=WD_TABLE_ALIGNMENT.LEFT
    tp=table._tbl.tblPr
    w=tp.find(qn("w:tblW"))
    if w is None: w=OxmlElement("w:tblW"); tp.append(w)
    w.set(qn("w:w"),str(TW)); w.set(qn("w:type"),"dxa")
    ind=tp.find(qn("w:tblInd"))
    if ind is None: ind=OxmlElement("w:tblInd"); tp.append(ind)
    ind.set(qn("w:w"),str(TI)); ind.set(qn("w:type"),"dxa")
    g=table._tbl.tblGrid
    for c in list(g): g.remove(c)
    for wd in widths:
        col=OxmlElement("w:gridCol"); col.set(qn("w:w"),str(wd)); g.append(col)
    for row in table.rows:
        for i,cell in enumerate(row.cells):
            tp=cell._tc.get_or_add_tcPr(); tcw=tp.find(qn("w:tcW"))
            if tcw is None: tcw=OxmlElement("w:tcW"); tp.append(tcw)
            tcw.set(qn("w:w"),str(widths[i])); tcw.set(qn("w:type"),"dxa")
            cell.width=Inches(widths[i]/1440); cell.vertical_alignment=WD_ALIGN_VERTICAL.CENTER
            cmargins(cell)

def rhead(row):
    tr=row._tr.get_or_add_trPr(); h=OxmlElement("w:tblHeader"); h.set(qn("w:val"),"true"); tr.append(h)

def font(run,name="Calibri",size=11,bold=None,italic=None,color=DARK):
    run.font.name=name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"),name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"),name)
    run.font.size=Pt(size); run.font.color.rgb=RGBColor.from_string(color)
    if bold is not None: run.bold=bold
    if italic is not None: run.italic=italic

def pborder(p,color=BLUE,size=10,space=4):
    pp=p._p.get_or_add_pPr(); bd=pp.find(qn("w:pBdr"))
    if bd is None: bd=OxmlElement("w:pBdr"); pp.append(bd)
    b=OxmlElement("w:bottom")
    b.set(qn("w:val"),"single"); b.set(qn("w:sz"),str(size)); b.set(qn("w:space"),str(space)); b.set(qn("w:color"),color)
    bd.append(b)

def pshade(p,fill):
    pp=p._p.get_or_add_pPr(); shd=pp.find(qn("w:shd"))
    if shd is None: shd=OxmlElement("w:shd"); pp.append(shd)
    shd.set(qn("w:fill"),fill)

def ctext(cell,text,bold=False,color=DARK,size=9.4,align=WD_ALIGN_PARAGRAPH.LEFT):
    p=cell.paragraphs[0]; p.alignment=align
    p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.1
    font(p.add_run(text),size=size,bold=bold,color=color)

def table(doc,headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers)); t.style="Table Grid"; geom(t,widths)
    h=t.rows[0]; rhead(h)
    for i,x in enumerate(headers):
        cshade(h.cells[i],LIGHT_BLUE); ctext(h.cells[i],x,bold=True,color=BLUE,size=9.2)
    for r in rows:
        row=t.add_row()
        for i,x in enumerate(r): ctext(row.cells[i],str(x),size=9.2)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

def code(doc,text):
    p=doc.add_paragraph(style="Code Block"); p.paragraph_format.keep_together=True; pshade(p,CODE_FILL)
    font(p.add_run(text.strip("\n")),name="Consolas",size=8.2,color="202124")

def callout(doc,label,text,kind="note"):
    fill={"note":LIGHT_BLUE,"warning":CAUTION_FILL,"risk":RISK_FILL,"success":GREEN_FILL}[kind]
    p=doc.add_paragraph(style="Callout"); pshade(p,fill)
    font(p.add_run(f"{label}: "),size=10.2,bold=True,color=BLUE if kind!="risk" else "9B1C1C")
    font(p.add_run(text),size=10.2,color=DARK)

def bullet(doc,text):
    p=doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent=Inches(0.375); p.paragraph_format.first_line_indent=Inches(-0.188)
    p.paragraph_format.space_after=Pt(3); p.paragraph_format.line_spacing=1.2
    font(p.add_run(text),size=10.5)

def num(doc,text,n):
    p=doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent=Inches(0.375); p.paragraph_format.first_line_indent=Inches(-0.188)
    p.paragraph_format.space_after=Pt(3); p.paragraph_format.line_spacing=1.2
    font(p.add_run(f"{n}.  {text}"),size=10.5)

def body(doc,text):
    p=doc.add_paragraph(style="Normal"); font(p.add_run(text))

def h(doc,text,level=1):
    p=doc.add_paragraph(text,style=f"Heading {level}"); p.paragraph_format.keep_with_next=True

def pageno(p):
    p.alignment=WD_ALIGN_PARAGRAPH.RIGHT; font(p.add_run("Page "),size=9,color=MUTED)
    for t,v in (("begin",None),("instr","PAGE"),("end",None)):
        if t=="instr": el=OxmlElement("w:instrText"); el.set(qn("xml:space"),"preserve"); el.text=v
        else: el=OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"),t)
        p.runs[0]._r.append(el)

def styles(doc):
    s=doc.styles; nm=s["Normal"]; nm.font.name="Calibri"
    nm._element.rPr.rFonts.set(qn("w:ascii"),"Calibri"); nm._element.rPr.rFonts.set(qn("w:hAnsi"),"Calibri")
    nm.font.size=Pt(11); nm.font.color.rgb=RGBColor.from_string(DARK)
    nm.paragraph_format.space_before=Pt(0); nm.paragraph_format.space_after=Pt(6); nm.paragraph_format.line_spacing=1.25
    for name,(sz,col,bf,af) in {"Heading 1":(16,BLUE2,16,9),"Heading 2":(13,BLUE2,12,6),"Heading 3":(12,BLUE,9,4)}.items():
        st=s[name]; st.font.name="Calibri"
        st._element.rPr.rFonts.set(qn("w:ascii"),"Calibri"); st._element.rPr.rFonts.set(qn("w:hAnsi"),"Calibri")
        st.font.size=Pt(sz); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(col)
        st.paragraph_format.space_before=Pt(bf); st.paragraph_format.space_after=Pt(af); st.paragraph_format.keep_with_next=True
    c=s.add_style("Code Block",1); c.font.name="Consolas"
    c._element.rPr.rFonts.set(qn("w:ascii"),"Consolas"); c._element.rPr.rFonts.set(qn("w:hAnsi"),"Consolas")
    c.font.size=Pt(8.2); c.paragraph_format.left_indent=Inches(0.12); c.paragraph_format.right_indent=Inches(0.12)
    c.paragraph_format.space_before=Pt(3); c.paragraph_format.space_after=Pt(6); c.paragraph_format.line_spacing=1.0
    co=s.add_style("Callout",1); co.font.name="Calibri"
    co._element.rPr.rFonts.set(qn("w:ascii"),"Calibri"); co._element.rPr.rFonts.set(qn("w:hAnsi"),"Calibri")
    co.font.size=Pt(10.2); co.paragraph_format.left_indent=Inches(0.12); co.paragraph_format.right_indent=Inches(0.12)
    co.paragraph_format.space_before=Pt(4); co.paragraph_format.space_after=Pt(8); co.paragraph_format.line_spacing=1.15

def section(doc,title):
    s=doc.sections[0]
    s.page_width=Inches(8.5); s.page_height=Inches(11)
    for m in ("top_margin","bottom_margin","left_margin","right_margin"): setattr(s,m,Inches(1))
    p=s.header.paragraphs[0]; p.paragraph_format.space_after=Pt(2)
    font(p.add_run(title),size=9,bold=True,color=MUTED); pborder(p,color="D0D5DD",size=4,space=3)
    fp=s.footer.paragraphs[0]; fp.paragraph_format.space_before=Pt(2); pageno(fp)

def cover(doc,tag,title,sub):
    for _ in range(4): doc.add_paragraph()
    for text,size,bold,color,after,it in [
        (tag,11,True,BLUE2,14,False),(title,26,True,BLUE,8,False),(sub,13,False,BLUE2,20,False),
        ("Prometheus + Grafana on RHEL 9",10.5,True,MUTED,0,False)]:
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        font(p.add_run(text),size=size,bold=bold,italic=it,color=color); p.paragraph_format.space_after=Pt(after)
    doc.add_page_break()


# =====================================================================
# DOC 1 — INTEGRATION
# =====================================================================
d=Document(); styles(d); section(d,"Integrate GitLab & Jenkins with Prometheus")
cover(d,"INTEGRATION GUIDE","Add GitLab & Jenkins to Monitoring","Target-side setup and Prometheus scrape config")

h(d,"Placeholders",1)
table(d,["Value","Meaning","Example"],[
    ("<MON_IP>","Monitoring server","10.10.10.11"),
    ("<GITLAB_IP>","GitLab server","10.10.10.20"),
    ("<GITLAB_FQDN>","GitLab URL host","gitlab.example.com"),
    ("<GIT_DATA_MOUNT>","Repo disk mount","/var/opt/gitlab"),
    ("<JENKINS_IP>","Jenkins server","10.10.10.21"),
    ("<JENKINS_PORT>","Jenkins port","8080"),
],[2400,3960,3000])

h(d,"1. GitLab",1)
h(d,"1.1 GitLab owners: expose the exporters",2)
body(d,"Edit /etc/gitlab/gitlab.rb, then run gitlab-ctl reconfigure.")
code(d,"""
gitlab_rails['monitoring_whitelist'] = ['127.0.0.0/8', '<MON_IP>/32']
node_exporter['listen_address']   = '<GITLAB_IP>:9100'
gitlab_exporter['listen_address'] = '<GITLAB_IP>'
gitlab_exporter['listen_port']    = '9168'
gitaly['configuration'] = { prometheus_listen_addr: '<GITLAB_IP>:9236' }
gitlab_workhorse['prometheus_listen_addr'] = '<GITLAB_IP>:9229'
sidekiq['metrics_enabled'] = true
sidekiq['listen_address']  = '<GITLAB_IP>'
sidekiq['listen_port']     = 8082
postgres_exporter['listen_address'] = '<GITLAB_IP>:9187'
redis_exporter['listen_address']    = '<GITLAB_IP>:9121'
""")
code(d,"""
sudo gitlab-ctl reconfigure
for p in 9100 9236 9168 9229 8082 9187 9121; do
  sudo firewall-cmd --permanent --add-rich-rule="rule family=\\"ipv4\\" \
    source address=\\"<MON_IP>/32\\" port protocol=\\"tcp\\" port=\\"$p\\" accept"
done
sudo firewall-cmd --reload
df -h /var/opt/gitlab/git-data      # note the mount for <GIT_DATA_MOUNT>
""")

h(d,"1.2 Monitoring server: add scrape jobs",2)
body(d,"Add under scrape_configs: in /etc/prometheus/prometheus.yml.")
code(d,"""
  - job_name: gitlab-node
    static_configs: [ { targets: ['<GITLAB_IP>:9100'], labels: { service: gitlab } } ]
  - job_name: gitlab-gitaly
    static_configs: [ { targets: ['<GITLAB_IP>:9236'], labels: { service: gitlab } } ]
  - job_name: gitlab-exporter
    static_configs: [ { targets: ['<GITLAB_IP>:9168'], labels: { service: gitlab } } ]
  - job_name: gitlab-workhorse
    static_configs: [ { targets: ['<GITLAB_IP>:9229'], labels: { service: gitlab } } ]
  - job_name: gitlab-sidekiq
    static_configs: [ { targets: ['<GITLAB_IP>:8082'], labels: { service: gitlab } } ]
  - job_name: gitlab-blackbox-https
    metrics_path: /probe
    params: { module: [http_2xx] }
    static_configs: [ { targets: ['https://<GITLAB_FQDN>/-/readiness'] } ]
    relabel_configs:
      - { source_labels: [__address__], target_label: __param_target }
      - { source_labels: [__param_target], target_label: instance }
      - { target_label: __address__, replacement: 127.0.0.1:9115 }
""")

h(d,"2. Jenkins",1)
h(d,"2.1 Jenkins host: node exporter",2)
body(d,"Install Node Exporter (platform runbook Section 5), then open 9100 to the monitoring server.")
code(d,"""
sudo firewall-cmd --permanent --add-rich-rule='rule family="ipv4" \
  source address="<MON_IP>/32" port protocol="tcp" port="9100" accept'
sudo firewall-cmd --reload
df -h /var/lib/jenkins           # note the mount for the JENKINS_HOME disk alert
""")

h(d,"2.2 Jenkins admin: install and expose metrics",2)
num(d,"Manage Jenkins -> Plugins -> Available -> install 'Prometheus metrics' -> restart.",1)
num(d,"Manage Jenkins -> System -> Prometheus section -> Path = prometheus -> Save.",2)
num(d,"Confirm: curl http://<JENKINS_IP>:<JENKINS_PORT>/prometheus/ returns metrics.",3)
num(d,"Restrict the Jenkins port to <MON_IP> at the firewall (or use an API token).",4)
callout(d,"Trailing slash","The endpoint is /prometheus/ with the slash. Prometheus must use metrics_path: /prometheus/ or it gets a 404.","note")

h(d,"2.3 Monitoring server: add scrape jobs",2)
code(d,"""
  - job_name: jenkins
    metrics_path: /prometheus/
    static_configs:
      - targets: ['<JENKINS_IP>:<JENKINS_PORT>']
        labels: { service: jenkins }
  - job_name: jenkins-blackbox
    metrics_path: /probe
    params: { module: [http_2xx] }
    static_configs: [ { targets: ['http://<JENKINS_IP>:<JENKINS_PORT>/login'] } ]
    relabel_configs:
      - { source_labels: [__address__], target_label: __param_target }
      - { source_labels: [__param_target], target_label: instance }
      - { target_label: __address__, replacement: 127.0.0.1:9115 }
""")

h(d,"3. Apply and verify",1)
code(d,"""
sudo promtool check config /etc/prometheus/prometheus.yml
sudo systemctl reload-or-restart prometheus
curl -s 'http://127.0.0.1:9090/api/v1/query?query=up{service="gitlab"}'
curl -s 'http://127.0.0.1:9090/api/v1/query?query=up{job="jenkins"}'
""")
body(d,"In Prometheus (Status -> Targets), every gitlab-* and jenkins target should read UP. If one is DOWN, curl the exporter port from the monitoring server and check the firewall on the target.")

d.core_properties.title="Integrate GitLab and Jenkins with Prometheus"
d.save(DIRP+"Integrate_GitLab_Jenkins.docx")
print(DIRP+"Integrate_GitLab_Jenkins.docx")


# =====================================================================
# DOC 2 — DASHBOARDS
# =====================================================================
d=Document(); styles(d); section(d,"Grafana Dashboards — GitLab & Jenkins")
cover(d,"DASHBOARD GUIDE","Grafana Dashboards","Import ready-made dashboards and build simple panels")

h(d,"1. Import a dashboard",1)
num(d,"In Grafana: Dashboards -> New -> Import.",1)
num(d,"Type the dashboard ID (from the table below) and click Load.",2)
num(d,"Pick your Prometheus data source from the dropdown.",3)
num(d,"Choose the folder (GitLab or Jenkins) and click Import.",4)
callout(d,"Data source","Every dashboard here uses the Prometheus data source. Do not use the GitLab data-source plugin - that queries the GitLab API, not your metrics.","note")

h(d,"2. Recommended dashboards",1)
body(d,"These are public dashboards on grafana.com. Import by ID.")
table(d,["Dashboard","ID","Use for"],[
    ("Node Exporter Full","1860","Host CPU / memory / disk - both servers"),
    ("GitLab Omnibus","22570","GitLab overall service health"),
    ("GitLab Self-managed - Gitaly","18917","Git RPC rate, latency, errors (git engine)"),
    ("Jenkins: Performance and Health","9964","Queue, executors, JVM heap, build results"),
    ("Prometheus Blackbox Exporter","7587","HTTPS / SSH front-door reachability"),
],[3400,900,5060])
callout(d,"Check versions","Dashboard panels depend on metric names that shift between GitLab and plugin versions. After import, open each panel; if one shows 'No data', edit its query to match a metric you actually have (Prometheus -> Graph tab).","warning")

h(d,"3. Build a simple panel yourself",1)
body(d,"If you want a small custom dashboard instead of the big imports: New -> Dashboard -> Add visualization -> pick Prometheus -> paste a query below -> Save.")
h(d,"GitLab panels",2)
bullet(d,"Component up/down (Stat):  up{service=\"gitlab\"}")
bullet(d,"Repo disk free % (Gauge):  100 * node_filesystem_avail_bytes{service=\"gitlab\",mountpoint=\"<GIT_DATA_MOUNT>\"} / node_filesystem_size_bytes{service=\"gitlab\",mountpoint=\"<GIT_DATA_MOUNT>\"}")
bullet(d,"Git RPC rate (Time series):  sum(rate(grpc_server_handled_total{job=\"gitlab-gitaly\"}[5m])) by (grpc_method)")
bullet(d,"Front door up (Stat):  probe_success{job=~\"gitlab-blackbox.*\"}")
h(d,"Jenkins panels",2)
bullet(d,"Jenkins up (Stat):  up{job=\"jenkins\"}")
bullet(d,"Build queue (Time series):  jenkins_queue_size_value")
bullet(d,"Executor use (Gauge):  jenkins_executor_in_use_value / jenkins_executor_count_value")
bullet(d,"Health score (Stat):  jenkins_health_check_score")
bullet(d,"Login reachable (Stat):  probe_success{job=\"jenkins-blackbox\"}")

h(d,"4. Tips",1)
bullet(d,"Put GitLab and Jenkins dashboards in their own folders so the list stays readable.")
bullet(d,"Set a dashboard variable for instance if you add more servers later.")
bullet(d,"Star the two or three dashboards you actually watch so they surface first.")

d.core_properties.title="Grafana Dashboards for GitLab and Jenkins"
d.save(DIRP+"Grafana_Dashboards_GitLab_Jenkins.docx")
print(DIRP+"Grafana_Dashboards_GitLab_Jenkins.docx")
