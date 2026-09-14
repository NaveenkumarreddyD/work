from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("/Users/naveendubba/Documents/Ibm mas and manage/Grafana_Production_Installation_RHEL9.docx")

BLUE = "1F4E78"
BLUE2 = "2E74B5"
DARK = "1F2937"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CODE_FILL = "F5F7FA"
CAUTION_FILL = "FFF4CE"
RISK_FILL = "FDECEC"
GREEN_FILL = "EAF5EA"
WHITE = "FFFFFF"
TABLE_WIDTH = 9360
TABLE_INDENT = 120


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    assert sum(widths) == TABLE_WIDTH
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths[idx] / 1440)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Calibri", size=11, bold=None, italic=None, color=DARK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_para_border_bottom(paragraph, color=BLUE, size=10, space=4):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def shade_paragraph(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=DARK, size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, text in enumerate(headers):
        set_cell_shading(header.cells[idx], LIGHT_BLUE)
        set_cell_text(header.cells[idx], text, bold=True, color=BLUE, size=9.2)
    for row_data in rows:
        row = table.add_row()
        for idx, text in enumerate(row_data):
            set_cell_text(row.cells[idx], str(text), size=9.2)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    return table


def add_code(doc, text):
    p = doc.add_paragraph(style="Code Block")
    p.paragraph_format.keep_together = True
    shade_paragraph(p, CODE_FILL)
    r = p.add_run(text.strip("\n"))
    set_run_font(r, name="Consolas", size=8.2, color="202124")
    return p


def add_callout(doc, label, text, kind="note"):
    fill = {"note": LIGHT_BLUE, "warning": CAUTION_FILL, "risk": RISK_FILL, "success": GREEN_FILL}[kind]
    p = doc.add_paragraph(style="Callout")
    shade_paragraph(p, fill)
    r = p.add_run(f"{label}: ")
    set_run_font(r, size=10.2, bold=True, color=BLUE if kind != "risk" else "9B1C1C")
    r = p.add_run(text)
    set_run_font(r, size=10.2, color=DARK)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Inches(0.375 + level * 0.25)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    set_run_font(r, size=10.5)
    return p


def add_number(doc, text, number):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(f"{number}.  {text}")
    set_run_font(r, size=10.5)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph(style="Normal")
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_run_font(r, bold=True)
        r = p.add_run(text[len(bold_lead):])
        set_run_font(r)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    tokens = {
        "Heading 1": (16, BLUE2, 18, 10),
        "Heading 2": (13, BLUE2, 14, 7),
        "Heading 3": (12, BLUE, 10, 5),
    }
    for name, (size, color, before, after) in tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    code = styles.add_style("Code Block", 1)
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code.font.size = Pt(8.2)
    code.paragraph_format.left_indent = Inches(0.12)
    code.paragraph_format.right_indent = Inches(0.12)
    code.paragraph_format.space_before = Pt(3)
    code.paragraph_format.space_after = Pt(6)
    code.paragraph_format.line_spacing = 1.0

    callout = styles.add_style("Callout", 1)
    callout.font.name = "Calibri"
    callout._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    callout._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    callout.font.size = Pt(10.2)
    callout.paragraph_format.left_indent = Inches(0.12)
    callout.paragraph_format.right_indent = Inches(0.12)
    callout.paragraph_format.space_before = Pt(4)
    callout.paragraph_format.space_after = Pt(8)
    callout.paragraph_format.line_spacing = 1.15


def configure_section(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Operations Runbook  |  Grafana Monitoring Platform")
    set_run_font(r, size=9, bold=True, color=MUTED)
    set_para_border_bottom(p, color="D0D5DD", size=4, space=3)

    footer = section.footer
    p = footer.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    add_page_number(p)


def page_break(doc):
    doc.add_page_break()


doc = Document()
configure_styles(doc)
for section in doc.sections:
    configure_section(section)

# Cover
for _ in range(5):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("PRODUCTION INSTALLATION RUNBOOK")
set_run_font(r, size=11, bold=True, color=BLUE2)
p.paragraph_format.space_after = Pt(18)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Grafana Monitoring Platform")
set_run_font(r, size=30, bold=True, color=BLUE)
p.paragraph_format.space_after = Pt(8)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Two-Server Deployment on Red Hat Enterprise Linux 9")
set_run_font(r, size=15, color=BLUE2)
p.paragraph_format.space_after = Pt(28)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Grafana OSS • PostgreSQL 16 • Prometheus LTS • Alertmanager • Node Exporter")
set_run_font(r, size=10.5, bold=True, color=MUTED)
p.paragraph_format.space_after = Pt(72)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Prepared for on-premises DevOps application monitoring")
set_run_font(r, size=10.5, italic=True, color=MUTED)
p.paragraph_format.space_after = Pt(6)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Version 1.0  |  23 July 2026")
set_run_font(r, size=10, color=MUTED)

page_break(doc)

add_heading(doc, "Document purpose and deployment decision", 1)
add_body(doc, "This runbook provides a complete build, configuration, validation, backup, failover, and recovery procedure for a monitoring platform deployed on two RHEL 9 servers. It covers Grafana OSS, PostgreSQL, Prometheus, Alertmanager, Node Exporter, HAProxy, and NGINX.")

add_callout(
    doc,
    "Important architectural constraint",
    "Two servers cannot provide safe automatic database quorum. This design therefore uses PostgreSQL primary/streaming-standby replication with controlled manual promotion. Prometheus and Alertmanager operate redundantly on both servers. Grafana is active on Server 1 and warm standby on Server 2. This is a production-ready active/standby design with documented recovery, not automatic active-active Grafana HA.",
    "warning",
)

add_table(
    doc,
    ["Item", "Selected design"],
    [
        ("Operating system", "Red Hat Enterprise Linux 9, x86_64"),
        ("Grafana", "Grafana OSS; active on Server 1, warm standby on Server 2"),
        ("Grafana database", "PostgreSQL 16 primary on Server 1, streaming standby on Server 2"),
        ("Metrics", "Prometheus 3.5.5 LTS on both servers"),
        ("Notifications", "Alertmanager 0.33.1 cluster on both servers"),
        ("Host metrics", "Node Exporter 1.12.1 on monitoring and application servers"),
        ("Query failover", "Local HAProxy selects the healthy Prometheus replica"),
        ("User access", "HTTPS through NGINX and a shared Grafana DNS name"),
        ("Database failover", "Manual promotion after fencing or isolating the failed primary"),
    ],
    [2500, 6860],
)

add_heading(doc, "Scope", 2)
for item in [
    "Installation and service configuration on both monitoring servers.",
    "PostgreSQL database initialization, replication, backup, promotion, and re-seeding.",
    "Prometheus collection, retention controls, alert rules, and replica validation.",
    "Alertmanager clustering and notification integration placeholders.",
    "Grafana, Prometheus data-source provisioning, NGINX HTTPS, and standby recovery.",
    "Integration patterns for GitLab, Jenkins, SonarQube, Nexus, and Linux servers.",
    "Firewall, SELinux, authentication, backup, patching, monitoring, and acceptance tests.",
]:
    add_bullet(doc, item)

add_heading(doc, "Out of scope", 2)
for item in [
    "Automatic PostgreSQL failover without a third quorum member.",
    "A long-term deduplicating metrics store such as Thanos or Grafana Mimir.",
    "Organization-specific SSO, SMTP, Teams, ServiceNow, or PagerDuty credentials.",
    "Application changes that require approval from GitLab, Jenkins, SonarQube, or Nexus owners.",
]:
    add_bullet(doc, item)

add_heading(doc, "Source-controlled artifacts", 2)
add_body(doc, "Store the following files in an approved private configuration repository. Never commit passwords, API tokens, private keys, or unencrypted database credentials.")
for item in [
    "/etc/prometheus/prometheus.yml and /etc/prometheus/rules/*.yml",
    "/etc/alertmanager/alertmanager.yml with secrets supplied at deployment time",
    "/etc/grafana/grafana.ini and /etc/grafana/provisioning/",
    "/etc/haproxy/haproxy.cfg and /etc/nginx/conf.d/grafana.conf",
    "Systemd unit files and the tested recovery procedure",
]:
    add_bullet(doc, item)

page_break(doc)

add_heading(doc, "1. Architecture and naming", 1)
add_heading(doc, "1.1 Component placement", 2)
add_table(
    doc,
    ["Component", "Server 1", "Server 2", "Operating mode"],
    [
        ("PostgreSQL 16", "Primary", "Streaming standby", "Manual promotion"),
        ("Grafana OSS", "Active", "Warm standby", "Start standby after DB promotion"),
        ("Prometheus LTS", "Replica A", "Replica B", "Both scrape all targets"),
        ("Alertmanager", "Replica A", "Replica B", "Clustered; both receive alerts"),
        ("Node Exporter", "Active", "Active", "Scraped by both Prometheus replicas"),
        ("HAProxy", "Active", "Active", "Local Prometheus primary/backup endpoint"),
        ("NGINX", "Active", "Installed/stopped", "HTTPS endpoint follows active Grafana"),
    ],
    [2100, 1900, 1900, 3460],
)

add_heading(doc, "1.2 Replace these values", 2)
add_callout(doc, "Before execution", "Replace every angle-bracket value in this runbook. Do not paste a command containing <PLACEHOLDER> into a production shell.", "risk")
add_table(
    doc,
    ["Placeholder", "Example", "Purpose"],
    [
        ("<MON1_FQDN>", "mon01.example.com", "Monitoring Server 1 hostname"),
        ("<MON1_IP>", "10.10.10.11", "Monitoring Server 1 address"),
        ("<MON2_FQDN>", "mon02.example.com", "Monitoring Server 2 hostname"),
        ("<MON2_IP>", "10.10.10.12", "Monitoring Server 2 address"),
        ("<GRAFANA_FQDN>", "grafana.example.com", "User-facing Grafana DNS"),
        ("<APPROVED_CIDR>", "10.20.0.0/16", "Authorized user network"),
        ("<PROM_RETENTION_SIZE>", "80GB", "At most 80% of dedicated Prometheus disk"),
        ("<SMTP_OR_WEBHOOK>", "Organization supplied", "Alert delivery destination"),
        ("<BACKUP_PATH>", "/backup/monitoring", "Approved protected backup destination"),
    ],
    [2200, 2600, 4560],
)

add_heading(doc, "1.3 Data flow", 2)
add_code(doc, """
Users --HTTPS/443--> Active NGINX --> Active Grafana
                                      |
                                      +--> Local HAProxy/9091
                                             |--> Prometheus A/9090
                                             +--> Prometheus B/9090

Prometheus A and B --> Node Exporter and application metrics endpoints
Prometheus A and B --> Alertmanager A and B --> Approved notification channel

Grafana Active --> PostgreSQL Primary
PostgreSQL Primary --> streaming replication --> PostgreSQL Standby
""")

add_heading(doc, "1.4 Availability behavior", 2)
add_table(
    doc,
    ["Failure", "Expected behavior", "Required action"],
    [
        ("One Prometheus replica", "Collection and dashboards continue through the surviving replica.", "Repair failed replica."),
        ("One Alertmanager replica", "Surviving member sends notifications.", "Repair failed member."),
        ("Server 2", "Primary Grafana and PostgreSQL continue; monitoring loses redundancy.", "Restore Server 2 promptly."),
        ("Server 1", "Server 2 continues metric collection and alerting. Grafana is unavailable until controlled DB promotion.", "Fence Server 1, promote PostgreSQL standby, start Grafana/NGINX, update DNS."),
        ("Network partition", "Alertmanager may intentionally send duplicate notifications. PostgreSQL standby must never be promoted unless the old primary is isolated.", "Follow incident command and fencing procedure."),
    ],
    [1700, 4300, 3360],
)

page_break(doc)

add_heading(doc, "2. Prerequisites and security approvals", 1)
add_heading(doc, "2.1 Server minimums", 2)
add_body(doc, "Final capacity must be calculated from target count, active time-series count, scrape interval, retention, dashboard concurrency, and alert volume. The following is a starting point, not an automatic sizing approval.")
add_table(
    doc,
    ["Resource", "Per server starting point", "Notes"],
    [
        ("CPU", "4-8 vCPU", "Prometheus query and ingestion load is the main driver."),
        ("Memory", "8-16 GB", "Increase for high series cardinality or concurrent dashboards."),
        ("OS disk", "40 GB", "Separate from Prometheus and PostgreSQL data where possible."),
        ("Prometheus data", "100 GB+ local SSD", "Do not use NFS. Size retention to 80% or less."),
        ("PostgreSQL data", "20 GB+ local SSD", "Low volume for Grafana but include WAL and backups."),
        ("Network", "Low-latency connection", "Both nodes scrape all application targets."),
    ],
    [1800, 2450, 5110],
)

add_heading(doc, "2.2 Required organizational inputs", 2)
for item in [
    "Two registered and patched RHEL 9 servers with sudo access.",
    "Forward and reverse DNS for both monitoring servers.",
    "Shared Grafana DNS name and a documented DNS failover procedure.",
    "TLS server certificate and private key for the shared Grafana DNS name.",
    "PostgreSQL server certificates for both database nodes and the issuing CA certificate.",
    "Approved notification service and credentials.",
    "Firewall approvals for the port matrix below.",
    "Approved backup location with access controls and retention policy.",
    "Application-owner approval to enable metrics endpoints or install plugins/exporters.",
]:
    add_bullet(doc, item)

add_heading(doc, "2.3 Port matrix", 2)
add_table(
    doc,
    ["Source", "Destination", "Port", "Purpose"],
    [
        ("Approved user networks", "Active Grafana/NGINX", "443/TCP", "Grafana HTTPS"),
        ("Both Grafana nodes", "Both Prometheus nodes", "9090/TCP", "Queries through HAProxy"),
        ("Both Prometheus nodes", "Both Alertmanagers", "9093/TCP", "Alert delivery"),
        ("Alertmanager nodes", "Peer Alertmanager", "9094/TCP+UDP", "Cluster gossip"),
        ("Both Prometheus nodes", "All Node Exporters", "9100/TCP", "Host metrics"),
        ("Both Prometheus nodes", "PostgreSQL exporters if added", "9187/TCP", "Optional DB metrics"),
        ("Server 2 PostgreSQL", "Server 1 PostgreSQL", "5432/TCP", "Streaming replication"),
        ("Both Grafana nodes", "Active PostgreSQL", "5432/TCP", "Grafana database"),
        ("Both Prometheus nodes", "Application endpoints", "Application-specific", "GitLab/Jenkins/SonarQube/Nexus metrics"),
    ],
    [2050, 2400, 1300, 3610],
)

add_callout(doc, "Security rule", "Do not expose Prometheus, Alertmanager, Node Exporter, or application metrics endpoints to public or general user networks. Restrict them to the two monitoring server IP addresses and approved administrators.", "warning")

add_heading(doc, "3. Base RHEL 9 preparation", 1)
add_heading(doc, "3.1 Run on both servers", 2)
add_code(doc, """
cat /etc/redhat-release
uname -m
hostname -f
df -h
free -h
timedatectl
chronyc tracking
""")

add_body(doc, "Confirm that `uname -m` returns `x86_64`. The binary URLs in this runbook are for AMD64/x86_64.")
add_code(doc, """
sudo dnf update -y
sudo dnf install -y curl wget tar gzip unzip vim policycoreutils-python-utils \
  nginx haproxy
sudo systemctl reboot
""")
add_callout(doc, "Change control", "Schedule and approve the RHEL patching and reboot. Reconnect after the server returns, then confirm the expected kernel and services.", "warning")

add_heading(doc, "3.2 Create data and backup mount points", 2)
add_body(doc, "Prefer dedicated local filesystems. The example retains the RHEL PostgreSQL default data location and creates an explicit backup directory.")
add_code(doc, """
sudo install -d -o root -g root -m 0750 /var/backups/monitoring
sudo install -d -o root -g root -m 0755 /etc/monitoring
df -hT /var/lib /var/backups/monitoring
""")

add_heading(doc, "3.3 Confirm SELinux and firewall state", 2)
add_code(doc, """
getenforce
sudo firewall-cmd --state
sudo firewall-cmd --get-active-zones
""")
add_body(doc, "Keep SELinux enforcing. Do not disable SELinux to solve a configuration error. Add only the required SELinux policy or boolean after confirming the denied operation.")

page_break(doc)

add_heading(doc, "4. Install PostgreSQL 16", 1)
add_callout(doc, "Database role", "Server 1 begins as PostgreSQL primary. Server 2 is initialized from Server 1 and remains a streaming standby until a controlled failover.", "note")

add_heading(doc, "4.1 Install packages on both servers", 2)
add_code(doc, """
sudo dnf module reset postgresql -y
sudo dnf module install postgresql:16/server -y
postgres --version
""")

add_heading(doc, "4.2 Initialize Server 1 only", 2)
add_code(doc, """
sudo PGSETUP_INITDB_OPTIONS="--data-checksums" postgresql-setup --initdb
sudo -u postgres pg_controldata /var/lib/pgsql/data | grep "Data page checksum"
""")
add_body(doc, "Expected result: `Data page checksum version` is non-zero.")

add_heading(doc, "4.3 Install PostgreSQL TLS files on Server 1", 2)
add_body(doc, "Copy the organization-issued Server 1 PostgreSQL certificate, private key, and issuing CA to the paths below.")
add_code(doc, """
sudo install -o postgres -g postgres -m 0600 <MON1_DB_PRIVATE_KEY> \
  /var/lib/pgsql/data/server.key
sudo install -o postgres -g postgres -m 0644 <MON1_DB_CERTIFICATE> \
  /var/lib/pgsql/data/server.crt
sudo install -o postgres -g postgres -m 0644 <DATABASE_CA_CERTIFICATE> \
  /var/lib/pgsql/data/root.crt
sudo restorecon -RFv /var/lib/pgsql/data
""")

add_heading(doc, "4.4 Configure Server 1", 2)
add_body(doc, "Edit `/var/lib/pgsql/data/postgresql.conf` and set:")
add_code(doc, """
listen_addresses = 'localhost,<MON1_IP>'
port = 5432
password_encryption = 'scram-sha-256'
ssl = on
ssl_cert_file = 'server.crt'
ssl_key_file = 'server.key'
ssl_ca_file = 'root.crt'

wal_level = replica
max_wal_senders = 10
max_replication_slots = 10
wal_keep_size = '2GB'
hot_standby = on
max_connections = 200
shared_buffers = '512MB'
log_connections = on
log_disconnections = on
log_line_prefix = '%m [%p] %u@%d %r '
""")

add_body(doc, "Append the following to `/var/lib/pgsql/data/pg_hba.conf` after replacing the addresses:")
add_code(doc, """
hostssl  grafana      grafana       <MON1_IP>/32  scram-sha-256
hostssl  grafana      grafana       <MON2_IP>/32  scram-sha-256
hostssl  replication  replicator    <MON2_IP>/32  scram-sha-256
""")

add_heading(doc, "4.5 Start Server 1 and create roles", 2)
add_code(doc, """
sudo systemctl enable --now postgresql
sudo systemctl status postgresql --no-pager
sudo -u postgres psql -c "SELECT version();"
""")

add_body(doc, "Open a PostgreSQL administrator session:")
add_code(doc, """
sudo -u postgres psql
""")
add_body(doc, "At the `psql` prompt, run:")
add_code(doc, """
CREATE ROLE grafana LOGIN;
\\password grafana

CREATE DATABASE grafana
  OWNER grafana
  ENCODING 'UTF8'
  TEMPLATE template0;

CREATE ROLE replicator WITH REPLICATION LOGIN;
\\password replicator

\\q
""")
add_callout(doc, "Password handling", "Use strong, unique passwords supplied by the approved secret-management process. Do not place passwords directly in shell commands, tickets, source control, or this document.", "warning")

add_heading(doc, "4.6 Allow PostgreSQL through the firewall", 2)
add_body(doc, "On Server 1, allow TCP 5432 only from the two monitoring addresses. Replace the example addresses:")
add_code(doc, """
sudo firewall-cmd --permanent --add-rich-rule='rule family="ipv4" \
  source address="<MON1_IP>/32" port protocol="tcp" port="5432" accept'
sudo firewall-cmd --permanent --add-rich-rule='rule family="ipv4" \
  source address="<MON2_IP>/32" port protocol="tcp" port="5432" accept'
sudo firewall-cmd --reload
""")

page_break(doc)

add_heading(doc, "5. Build the PostgreSQL streaming standby", 1)
add_heading(doc, "5.1 Prepare Server 2", 2)
add_body(doc, "Do not run `postgresql-setup --initdb` on Server 2. Stop PostgreSQL and confirm the target directory is empty before taking the base backup.")
add_code(doc, """
sudo systemctl stop postgresql
sudo install -d -o postgres -g postgres -m 0700 /var/lib/pgsql/data
sudo find /var/lib/pgsql/data -mindepth 1 -maxdepth 1 -print
""")
add_callout(doc, "Destructive check", "If the command lists files, stop. Confirm that Server 2 contains no required PostgreSQL data before clearing or replacing the directory.", "risk")

add_heading(doc, "5.2 Create the replication password file", 2)
add_code(doc, """
sudo -u postgres vi /var/lib/pgsql/.pgpass
""")
add_body(doc, "Add one line:")
add_code(doc, """
<MON1_FQDN>:5432:replication:replicator:<REPLICATION_PASSWORD>
""")
add_code(doc, """
sudo chown postgres:postgres /var/lib/pgsql/.pgpass
sudo chmod 0600 /var/lib/pgsql/.pgpass
""")

add_heading(doc, "5.3 Take the base backup", 2)
add_code(doc, """
sudo -u postgres pg_basebackup \
  --host=<MON1_FQDN> \
  --port=5432 \
  --username=replicator \
  --pgdata=/var/lib/pgsql/data \
  --format=plain \
  --wal-method=stream \
  --progress \
  --write-recovery-conf \
  --create-slot \
  --slot=grafana_standby
""")

add_heading(doc, "5.4 Replace Server 2 TLS identity", 2)
add_body(doc, "The base backup copies Server 1 configuration files. Replace the copied certificate and key with Server 2’s identity before starting PostgreSQL.")
add_code(doc, """
sudo install -o postgres -g postgres -m 0600 <MON2_DB_PRIVATE_KEY> \
  /var/lib/pgsql/data/server.key
sudo install -o postgres -g postgres -m 0644 <MON2_DB_CERTIFICATE> \
  /var/lib/pgsql/data/server.crt
sudo install -o postgres -g postgres -m 0644 <DATABASE_CA_CERTIFICATE> \
  /var/lib/pgsql/data/root.crt
sudo restorecon -RFv /var/lib/pgsql/data
""")
add_body(doc, "Edit Server 2 `/var/lib/pgsql/data/postgresql.conf` and change:")
add_code(doc, """
listen_addresses = 'localhost,<MON2_IP>'
""")

add_heading(doc, "5.5 Start and validate the standby", 2)
add_code(doc, """
sudo systemctl enable --now postgresql
sudo systemctl status postgresql --no-pager
sudo -u postgres psql -c "SELECT pg_is_in_recovery();"
""")
add_body(doc, "Expected on Server 2: `pg_is_in_recovery` returns `t`.")
add_body(doc, "On Server 1:")
add_code(doc, """
sudo -u postgres psql -x -c \
  "SELECT application_name, client_addr, state, sync_state, replay_lag \
   FROM pg_stat_replication;"
""")
add_body(doc, "Expected: one streaming replica. This design uses asynchronous replication to preserve Grafana availability if the standby is temporarily unavailable. Monitor replication lag and define the accepted recovery-point objective.")

add_heading(doc, "5.6 Verify read-only behavior", 2)
add_code(doc, """
sudo -u postgres psql -d grafana -c \
  "SELECT pg_is_in_recovery(), current_timestamp;"
""")

page_break(doc)

add_heading(doc, "6. Install Node Exporter", 1)
add_body(doc, "Install Node Exporter on both monitoring servers and later repeat the same procedure on every approved Linux application server.")
add_heading(doc, "6.1 Download and verify on each server", 2)
add_code(doc, """
cd /tmp
curl -fLO https://github.com/prometheus/node_exporter/releases/download/v1.12.1/\
node_exporter-1.12.1.linux-amd64.tar.gz

echo "b51d8a76aa2a9156a55d501aca6276fae09e262259a5e4e831d2c2222f084e63  \
node_exporter-1.12.1.linux-amd64.tar.gz" | sha256sum -c -

tar -xzf node_exporter-1.12.1.linux-amd64.tar.gz
sudo useradd --system --no-create-home --shell /sbin/nologin node_exporter
sudo install -m 0755 node_exporter-1.12.1.linux-amd64/node_exporter \
  /usr/local/bin/node_exporter
""")

add_heading(doc, "6.2 Create the service", 2)
add_code(doc, """
sudo tee /etc/systemd/system/node_exporter.service >/dev/null <<'EOF'
[Unit]
Description=Prometheus Node Exporter
Wants=network-online.target
After=network-online.target

[Service]
User=node_exporter
Group=node_exporter
Type=simple
ExecStart=/usr/local/bin/node_exporter --web.listen-address=0.0.0.0:9100
Restart=on-failure
RestartSec=5
NoNewPrivileges=true
PrivateTmp=true
ProtectHome=true
ProtectSystem=strict

[Install]
WantedBy=multi-user.target
EOF

sudo systemctl daemon-reload
sudo systemctl enable --now node_exporter
sudo systemctl status node_exporter --no-pager
curl -fsS http://localhost:9100/metrics | head
""")

add_heading(doc, "6.3 Firewall", 2)
add_body(doc, "On each Node Exporter host, allow TCP 9100 only from both Prometheus addresses.")
add_code(doc, """
sudo firewall-cmd --permanent --add-rich-rule='rule family="ipv4" \
  source address="<MON1_IP>/32" port protocol="tcp" port="9100" accept'
sudo firewall-cmd --permanent --add-rich-rule='rule family="ipv4" \
  source address="<MON2_IP>/32" port protocol="tcp" port="9100" accept'
sudo firewall-cmd --reload
""")

page_break(doc)

add_heading(doc, "7. Install Prometheus LTS on both servers", 1)
add_heading(doc, "7.1 Download and verify", 2)
add_code(doc, """
cd /tmp
curl -fLO https://github.com/prometheus/prometheus/releases/download/v3.5.5/\
prometheus-3.5.5.linux-amd64.tar.gz

echo "64d0beab873272b861a91df41668bc852c7e2e5b23f75c16059fb15b5630c577  \
prometheus-3.5.5.linux-amd64.tar.gz" | sha256sum -c -

tar -xzf prometheus-3.5.5.linux-amd64.tar.gz
sudo useradd --system --no-create-home --shell /sbin/nologin prometheus
sudo install -d -o prometheus -g prometheus -m 0750 \
  /etc/prometheus /etc/prometheus/rules /var/lib/prometheus
sudo install -m 0755 prometheus-3.5.5.linux-amd64/prometheus \
  /usr/local/bin/prometheus
sudo install -m 0755 prometheus-3.5.5.linux-amd64/promtool \
  /usr/local/bin/promtool
""")

add_heading(doc, "7.2 Prometheus configuration", 2)
add_body(doc, "Create `/etc/prometheus/prometheus.yml` on both servers. On Server 2 change the replica label from `prometheus-a` to `prometheus-b`.")
add_code(doc, """
global:
  scrape_interval: 30s
  scrape_timeout: 10s
  evaluation_interval: 30s
  external_labels:
    environment: production
    replica: prometheus-a

rule_files:
  - /etc/prometheus/rules/*.yml

alerting:
  alertmanagers:
    - static_configs:
        - targets:
            - <MON1_IP>:9093
            - <MON2_IP>:9093

scrape_configs:
  - job_name: prometheus
    static_configs:
      - targets:
          - <MON1_IP>:9090
          - <MON2_IP>:9090

  - job_name: monitoring-servers
    static_configs:
      - targets:
          - <MON1_IP>:9100
          - <MON2_IP>:9100

  - job_name: application-linux-servers
    static_configs:
      - targets:
          - <GITLAB_SERVER>:9100
          - <JENKINS_SERVER>:9100
          - <SONARQUBE_SERVER>:9100
          - <NEXUS_SERVER>:9100
""")

add_heading(doc, "7.3 Initial alert rules", 2)
add_body(doc, "Create `/etc/prometheus/rules/infrastructure.yml` on both servers:")
add_code(doc, """
groups:
  - name: infrastructure
    rules:
      - alert: TargetDown
        expr: up == 0
        for: 5m
        labels:
          severity: critical
        annotations:
          summary: "Monitoring target is unavailable"
          description: "{{ $labels.job }} target {{ $labels.instance }} is down."

      - alert: HostHighCPU
        expr: 100 - (avg by(instance)
          (rate(node_cpu_seconds_total{mode="idle"}[5m])) * 100) > 90
        for: 15m
        labels:
          severity: warning
        annotations:
          summary: "High CPU usage on {{ $labels.instance }}"

      - alert: HostLowMemory
        expr: (node_memory_MemAvailable_bytes /
          node_memory_MemTotal_bytes) * 100 < 10
        for: 10m
        labels:
          severity: warning
        annotations:
          summary: "Available memory below 10% on {{ $labels.instance }}"

      - alert: HostFilesystemLow
        expr: 100 * node_filesystem_avail_bytes{fstype!~"tmpfs|overlay"} /
          node_filesystem_size_bytes{fstype!~"tmpfs|overlay"} < 15
        for: 15m
        labels:
          severity: warning
        annotations:
          summary: "Filesystem space below 15% on {{ $labels.instance }}"
""")

add_heading(doc, "7.4 Validate ownership and syntax", 2)
add_code(doc, """
sudo chown -R prometheus:prometheus /etc/prometheus /var/lib/prometheus
sudo promtool check config /etc/prometheus/prometheus.yml
sudo promtool check rules /etc/prometheus/rules/infrastructure.yml
""")

add_heading(doc, "7.5 Create the Prometheus service", 2)
add_callout(doc, "Retention value", "Replace <PROM_RETENTION_SIZE> with no more than approximately 80% of the dedicated Prometheus filesystem. Example: use 80GB only when at least 100GB is allocated exclusively to Prometheus.", "warning")
add_code(doc, """
sudo tee /etc/systemd/system/prometheus.service >/dev/null <<'EOF'
[Unit]
Description=Prometheus Monitoring Server
Wants=network-online.target
After=network-online.target

[Service]
User=prometheus
Group=prometheus
Type=simple
ExecStart=/usr/local/bin/prometheus \
  --config.file=/etc/prometheus/prometheus.yml \
  --storage.tsdb.path=/var/lib/prometheus \
  --storage.tsdb.retention.time=30d \
  --storage.tsdb.retention.size=<PROM_RETENTION_SIZE> \
  --web.listen-address=0.0.0.0:9090
Restart=on-failure
RestartSec=5
LimitNOFILE=65536
NoNewPrivileges=true
PrivateTmp=true
ProtectHome=true
ProtectSystem=full
ReadWritePaths=/var/lib/prometheus

[Install]
WantedBy=multi-user.target
EOF

sudo systemctl daemon-reload
sudo systemctl enable --now prometheus
sudo systemctl status prometheus --no-pager
curl -fsS http://localhost:9090/-/healthy
curl -fsS http://localhost:9090/-/ready
""")

add_heading(doc, "7.6 Network access", 2)
add_body(doc, "Allow TCP 9090 between the two monitoring servers. Do not expose the Prometheus UI broadly.")

page_break(doc)

add_heading(doc, "8. Install Alertmanager on both servers", 1)
add_heading(doc, "8.1 Download and verify", 2)
add_code(doc, """
cd /tmp
curl -fLO https://github.com/prometheus/alertmanager/releases/download/v0.33.1/\
alertmanager-0.33.1.linux-amd64.tar.gz

echo "93d802cba6a8d27239d747ce117df7648d326ab67394e32247540b030e9842ba  \
alertmanager-0.33.1.linux-amd64.tar.gz" | sha256sum -c -

tar -xzf alertmanager-0.33.1.linux-amd64.tar.gz
sudo useradd --system --no-create-home --shell /sbin/nologin alertmanager
sudo install -d -o alertmanager -g alertmanager -m 0750 \
  /etc/alertmanager /var/lib/alertmanager
sudo install -m 0755 alertmanager-0.33.1.linux-amd64/alertmanager \
  /usr/local/bin/alertmanager
sudo install -m 0755 alertmanager-0.33.1.linux-amd64/amtool \
  /usr/local/bin/amtool
""")

add_heading(doc, "8.2 Configure routing", 2)
add_body(doc, "Create `/etc/alertmanager/alertmanager.yml` on both servers. The initial null receiver is valid for build validation but must be replaced before go-live.")
add_code(doc, """
global:
  resolve_timeout: 5m

route:
  receiver: operations-placeholder
  group_by: [alertname, job]
  group_wait: 30s
  group_interval: 5m
  repeat_interval: 4h

receivers:
  - name: operations-placeholder
""")
add_callout(doc, "Go-live gate", "Production acceptance must fail until a test alert is successfully delivered to the approved notification channel and acknowledged by the support team.", "risk")

add_heading(doc, "8.3 Validate configuration", 2)
add_code(doc, """
sudo chown -R alertmanager:alertmanager \
  /etc/alertmanager /var/lib/alertmanager
sudo amtool check-config /etc/alertmanager/alertmanager.yml
""")

add_heading(doc, "8.4 Server 1 service", 2)
add_code(doc, """
sudo tee /etc/systemd/system/alertmanager.service >/dev/null <<'EOF'
[Unit]
Description=Prometheus Alertmanager
Wants=network-online.target
After=network-online.target

[Service]
User=alertmanager
Group=alertmanager
Type=simple
ExecStart=/usr/local/bin/alertmanager \
  --config.file=/etc/alertmanager/alertmanager.yml \
  --storage.path=/var/lib/alertmanager \
  --web.listen-address=0.0.0.0:9093 \
  --cluster.listen-address=0.0.0.0:9094 \
  --cluster.advertise-address=<MON1_IP>:9094 \
  --cluster.peer=<MON2_IP>:9094
Restart=on-failure
RestartSec=5
NoNewPrivileges=true
PrivateTmp=true
ProtectHome=true
ProtectSystem=full
ReadWritePaths=/var/lib/alertmanager

[Install]
WantedBy=multi-user.target
EOF
""")

add_heading(doc, "8.5 Server 2 service", 2)
add_body(doc, "Use the same service, changing only:")
add_code(doc, """
--cluster.advertise-address=<MON2_IP>:9094
--cluster.peer=<MON1_IP>:9094
""")

add_heading(doc, "8.6 Start and validate", 2)
add_code(doc, """
sudo systemctl daemon-reload
sudo systemctl enable --now alertmanager
sudo systemctl status alertmanager --no-pager
curl -fsS http://localhost:9093/-/healthy
curl -fsS http://localhost:9093/-/ready
curl -fsS http://localhost:9093/api/v2/status
""")
add_body(doc, "Allow 9093/TCP from both Prometheus nodes. Allow 9094/TCP and UDP only between the two Alertmanager nodes.")

page_break(doc)

add_heading(doc, "9. Configure Prometheus query failover", 1)
add_body(doc, "Grafana queries a local HAProxy listener on `127.0.0.1:9091`. HAProxy uses the local Prometheus replica first and the remote replica as backup.")

add_heading(doc, "9.1 Server 1 HAProxy configuration", 2)
add_code(doc, """
global
    log /dev/log local0
    user haproxy
    group haproxy
    daemon

defaults
    log global
    mode http
    timeout connect 5s
    timeout client 60s
    timeout server 60s

frontend prometheus_query
    bind 127.0.0.1:9091
    default_backend prometheus_servers

backend prometheus_servers
    option httpchk GET /-/ready
    http-check expect status 200
    server prometheus-a <MON1_IP>:9090 check
    server prometheus-b <MON2_IP>:9090 check backup
""")

add_heading(doc, "9.2 Server 2 HAProxy configuration", 2)
add_body(doc, "Use the same file but reverse primary and backup:")
add_code(doc, """
server prometheus-b <MON2_IP>:9090 check
server prometheus-a <MON1_IP>:9090 check backup
""")

add_heading(doc, "9.3 Validate and start", 2)
add_code(doc, """
sudo haproxy -c -f /etc/haproxy/haproxy.cfg
sudo systemctl enable --now haproxy
sudo systemctl status haproxy --no-pager
curl -fsS http://127.0.0.1:9091/-/ready
""")

page_break(doc)

add_heading(doc, "10. Install Grafana OSS on both servers", 1)
add_heading(doc, "10.1 Add the official repository", 2)
add_code(doc, """
cd /tmp
wget -q -O grafana-gpg.key https://rpm.grafana.com/gpg.key
sudo rpm --import grafana-gpg.key

sudo tee /etc/yum.repos.d/grafana.repo >/dev/null <<'EOF'
[grafana]
name=grafana
baseurl=https://rpm.grafana.com
repo_gpgcheck=1
enabled=1
gpgcheck=1
gpgkey=https://rpm.grafana.com/gpg.key
sslverify=1
sslcacert=/etc/pki/tls/certs/ca-bundle.crt
EOF

sudo dnf makecache --refresh
sudo dnf install -y grafana
rpm -q grafana
""")

add_heading(doc, "10.2 Create shared Grafana secrets", 2)
add_body(doc, "The database password and Grafana `secret_key` must be identical on both servers. Store them through the organizational secret-management process.")
add_code(doc, """
sudo install -d -o root -g grafana -m 0750 /etc/grafana/secrets
sudo vi /etc/grafana/secrets/postgres_password
sudo vi /etc/grafana/secrets/secret_key
sudo chown root:grafana /etc/grafana/secrets/*
sudo chmod 0640 /etc/grafana/secrets/*
""")
add_body(doc, "Generate the secret key on an approved secure workstation, store it in the secret manager, and deploy the same value to both nodes. Do not regenerate it during failover.")

add_heading(doc, "10.3 Install the database CA", 2)
add_code(doc, """
sudo install -o root -g grafana -m 0640 <DATABASE_CA_CERTIFICATE> \
  /etc/grafana/postgresql-ca.crt
""")

add_heading(doc, "10.4 Configure Server 1 Grafana", 2)
add_body(doc, "Edit `/etc/grafana/grafana.ini`:")
add_code(doc, """
[database]
type = postgres
host = <MON1_FQDN>:5432
name = grafana
user = grafana
password = $__file{/etc/grafana/secrets/postgres_password}
ssl_mode = verify-full
ca_cert_path = /etc/grafana/postgresql-ca.crt
max_idle_conn = 10
max_open_conn = 100
conn_max_lifetime = 3600
migration_locking = true
locking_attempt_timeout_sec = 60

[server]
protocol = http
http_addr = 127.0.0.1
http_port = 3000
domain = <GRAFANA_FQDN>
root_url = https://<GRAFANA_FQDN>/

[security]
secret_key = $__file{/etc/grafana/secrets/secret_key}
cookie_secure = true
cookie_samesite = strict

[users]
allow_sign_up = false

[auth.anonymous]
enabled = false

[auth.basic]
enabled = true
password_policy = true

[unified_alerting]
enabled = false
""")

add_heading(doc, "10.5 Configure Server 2 Grafana", 2)
add_body(doc, "Use the same file. While Server 1 is primary, Server 2 also points to `<MON1_FQDN>:5432`. Keep Server 2 Grafana stopped until failover testing or recovery.")

add_heading(doc, "10.6 Provision Prometheus", 2)
add_code(doc, """
sudo tee /etc/grafana/provisioning/datasources/prometheus.yml >/dev/null <<'EOF'
apiVersion: 1

datasources:
  - name: Prometheus
    type: prometheus
    access: proxy
    url: http://127.0.0.1:9091
    isDefault: true
    editable: false
EOF
""")

add_heading(doc, "10.7 Start Grafana on Server 1", 2)
add_code(doc, """
sudo systemctl daemon-reload
sudo systemctl enable --now grafana-server
sudo systemctl status grafana-server --no-pager
curl -fsS http://127.0.0.1:3000/api/health
""")
add_body(doc, "Expected: database status is `ok`. Review logs if startup fails:")
add_code(doc, """
sudo journalctl -u grafana-server --since "30 minutes ago" --no-pager
""")

add_heading(doc, "10.8 Keep Server 2 in warm standby", 2)
add_code(doc, """
sudo systemctl disable --now grafana-server
""")

add_heading(doc, "11. Configure NGINX and HTTPS", 1)
add_heading(doc, "11.1 Install the shared Grafana certificate on both servers", 2)
add_code(doc, """
sudo install -o root -g root -m 0644 <GRAFANA_TLS_CERTIFICATE> \
  /etc/pki/tls/certs/grafana.crt
sudo install -o root -g root -m 0600 <GRAFANA_TLS_PRIVATE_KEY> \
  /etc/pki/tls/private/grafana.key
""")

add_heading(doc, "11.2 Create NGINX configuration on both servers", 2)
add_code(doc, """
sudo tee /etc/nginx/conf.d/grafana.conf >/dev/null <<'EOF'
map $http_upgrade $connection_upgrade {
    default upgrade;
    '' close;
}

server {
    listen 443 ssl;
    server_name <GRAFANA_FQDN>;

    ssl_certificate /etc/pki/tls/certs/grafana.crt;
    ssl_certificate_key /etc/pki/tls/private/grafana.key;
    ssl_protocols TLSv1.2 TLSv1.3;

    location / {
        proxy_set_header Host $host;
        proxy_set_header X-Real-IP $remote_addr;
        proxy_set_header X-Forwarded-For $proxy_add_x_forwarded_for;
        proxy_set_header X-Forwarded-Proto https;
        proxy_http_version 1.1;
        proxy_set_header Upgrade $http_upgrade;
        proxy_set_header Connection $connection_upgrade;
        proxy_pass http://127.0.0.1:3000;
    }
}
EOF

sudo setsebool -P httpd_can_network_connect 1
sudo nginx -t
""")

add_heading(doc, "11.3 Start NGINX on Server 1", 2)
add_code(doc, """
sudo systemctl enable --now nginx
sudo systemctl status nginx --no-pager
sudo firewall-cmd --permanent --add-service=https
sudo firewall-cmd --reload
""")
add_body(doc, "Keep NGINX installed but stopped on Server 2 until controlled failover:")
add_code(doc, """
sudo systemctl disable --now nginx
""")

add_heading(doc, "11.4 DNS", 2)
add_body(doc, "Point `<GRAFANA_FQDN>` to Server 1 for normal operations. Use a low but policy-compliant TTL to support controlled recovery. Document the owner and exact DNS change procedure.")

add_heading(doc, "11.5 Initial login", 2)
add_body(doc, "Open `https://<GRAFANA_FQDN>/`. Sign in with `admin/admin`, then immediately replace the default password. Configure approved users and roles. If SSO is required, complete and test the organization’s supported identity-provider integration before go-live.")

add_heading(doc, "12. Integrate application metrics", 1)
add_heading(doc, "12.1 Linux host metrics", 2)
add_body(doc, "Install Node Exporter on each GitLab, Jenkins, SonarQube, and Nexus Linux server. Restrict port 9100 to the two Prometheus IP addresses. Add each target to both Prometheus configurations.")

add_heading(doc, "12.2 GitLab Self-Managed", 2)
for item in [
    "Confirm the GitLab installation method and version.",
    "Use GitLab’s built-in Prometheus metrics and bundled exporters where supported.",
    "Allow only the two Prometheus server addresses in the monitoring allowlist.",
    "Do not expose unauthenticated exporters outside the monitoring network.",
    "Collect GitLab, Sidekiq, Gitaly, Workhorse, PostgreSQL, Redis, Registry, and node metrics as applicable.",
]:
    add_bullet(doc, item)

add_heading(doc, "12.3 Jenkins", 2)
for item in [
    "Install the Jenkins Prometheus Metrics plugin through the approved plugin process.",
    "Confirm the plugin supports the installed Jenkins controller version.",
    "Validate the default `/prometheus/` endpoint, including its required trailing slash.",
    "Use Jenkins authentication and network restrictions consistent with organizational policy.",
    "Validate job status, duration, executor, queue, JVM, and controller metrics before dashboard acceptance.",
]:
    add_bullet(doc, item)

add_heading(doc, "12.4 SonarQube", 2)
for item in [
    "Use `/api/monitoring/metrics` for supported operational metrics and secure it with the configured Sonar system passcode.",
    "Use `/api/measures` for project code-quality values and histories.",
    "Because code-quality API responses are not identical to ordinary Prometheus scraping, use an approved API exporter or collection process where required.",
    "Also collect CPU, memory, disk, and Java process metrics through Node Exporter and approved JMX collection.",
]:
    add_bullet(doc, item)

add_heading(doc, "12.5 Nexus Repository", 2)
for item in [
    "For Nexus Repository 3.81 or later, use `/service/rest/metrics/prometheus`.",
    "For earlier supported releases, validate `/service/metrics/prometheus`.",
    "Use a dedicated account with the `nx-metrics-all` privilege.",
    "Use synthetic HTTP checks for required artifact or repository availability because native instance metrics may not prove that a specific artifact can be downloaded.",
]:
    add_bullet(doc, item)

add_heading(doc, "12.6 Safe rollout procedure", 2)
for number, text in enumerate([
    "Enable one application endpoint at a time.",
    "Validate the endpoint locally on the application server.",
    "Validate network access from both Prometheus servers.",
    "Add the target to both Prometheus configurations.",
    "Run `promtool check config` on both nodes.",
    "Restart both Prometheus services during the approved window.",
    "Confirm the target is `UP` on both replicas.",
    "Validate metrics with the application owner before creating alerts.",
], start=1):
    add_number(doc, text, number)

add_heading(doc, "13. Backups and recovery", 1)
add_heading(doc, "13.1 What must be backed up", 2)
add_table(
    doc,
    ["Artifact", "Backup method", "Frequency"],
    [
        ("Grafana PostgreSQL database", "Nightly `pg_dump` plus retention in approved backup storage", "Daily and before upgrades"),
        ("PostgreSQL configuration", "Protected file backup", "After every approved change"),
        ("Grafana configuration and provisioning", "Protected backup and source control without secrets", "After every change"),
        ("Grafana plugins", "Package/plugin inventory; reinstall from approved source", "After plugin change"),
        ("Prometheus/Alertmanager configuration", "Private source control", "After every change"),
        ("Prometheus data", "Redundant replicas; optional snapshots if history requires backup", "Policy dependent"),
        ("TLS certificates and keys", "PKI/secret-management system", "According to PKI policy"),
    ],
    [2700, 4300, 2360],
)

add_heading(doc, "13.2 PostgreSQL logical backup", 2)
add_body(doc, "Run on Server 1 using an approved protected path:")
add_code(doc, """
sudo -u postgres pg_dump \
  --format=custom \
  --file=<BACKUP_PATH>/grafana_$(date +%F_%H%M).dump \
  grafana

sudo -u postgres pg_restore --list \
  <BACKUP_PATH>/<GRAFANA_BACKUP_FILE>.dump | head
""")
add_body(doc, "Copy backups off Server 1. Streaming replication is not a substitute for backup because accidental deletions and corruption can replicate.")

add_heading(doc, "13.3 Configuration backup", 2)
add_code(doc, """
sudo tar --xattrs --selinux -czf \
  <BACKUP_PATH>/monitoring-config_$(date +%F_%H%M).tgz \
  /etc/grafana \
  /etc/prometheus \
  /etc/alertmanager \
  /etc/haproxy \
  /etc/nginx/conf.d/grafana.conf \
  /etc/systemd/system/prometheus.service \
  /etc/systemd/system/alertmanager.service \
  /etc/systemd/system/node_exporter.service
""")
add_callout(doc, "Sensitive content", "The archive contains secrets and private configuration. Encrypt it, tightly restrict access, and never attach it to an ordinary ticket.", "warning")

add_heading(doc, "13.4 Restore test", 2)
for item in [
    "Perform a scheduled restore to a non-production PostgreSQL instance.",
    "Validate Grafana startup against the restored database.",
    "Confirm dashboards, users, data sources, and permissions.",
    "Record restoration time and compare it with the recovery-time objective.",
    "Correct the procedure after every failed or incomplete restore test.",
]:
    add_bullet(doc, item)

page_break(doc)

add_heading(doc, "14. Controlled failover to Server 2", 1)
add_callout(doc, "Critical safety requirement", "Never promote Server 2 while Server 1 PostgreSQL may still accept writes. First fence, power off, or network-isolate Server 1. This prevents split-brain and divergent Grafana databases.", "risk")

add_heading(doc, "14.1 Decision and fencing", 2)
for number, text in enumerate([
    "Declare the monitoring-platform incident and assign an incident commander.",
    "Confirm Server 1 is unavailable and not merely isolated from the operator.",
    "Fence or power off Server 1, or block its PostgreSQL and Grafana network access.",
    "Record the last known PostgreSQL replication lag and accepted data-loss risk.",
    "Approve promotion of Server 2.",
], start=1):
    add_number(doc, text, number)

add_heading(doc, "14.2 Promote PostgreSQL on Server 2", 2)
add_code(doc, """
sudo -u postgres psql -c "SELECT pg_is_in_recovery();"
sudo -u postgres psql -c "SELECT pg_promote(wait_seconds => 60);"
sudo -u postgres psql -c "SELECT pg_is_in_recovery();"
""")
add_body(doc, "Expected after promotion: `pg_is_in_recovery` returns `f`.")

add_heading(doc, "14.3 Point Server 2 Grafana to the new primary", 2)
add_body(doc, "Edit Server 2 `/etc/grafana/grafana.ini`:")
add_code(doc, """
[database]
host = <MON2_FQDN>:5432
""")
add_code(doc, """
sudo systemctl enable --now grafana-server
sudo systemctl enable --now nginx
sudo systemctl status grafana-server nginx --no-pager
curl -fsS http://127.0.0.1:3000/api/health
""")

add_heading(doc, "14.4 Update DNS", 2)
add_body(doc, "Change `<GRAFANA_FQDN>` to Server 2 only after local validation succeeds. Confirm HTTPS, login, dashboards, data sources, and permissions.")

add_heading(doc, "14.5 Rebuild Server 1 as standby", 2)
add_body(doc, "After Server 1 is repaired, do not simply restart its old PostgreSQL database. Re-seed it from the new primary using a fresh `pg_basebackup`, or use `pg_rewind` only under a separately tested DBA procedure. A fresh base backup is slower but operationally simpler and safer.")

add_heading(doc, "14.6 Return to normal roles", 2)
add_body(doc, "Fail back only during an approved maintenance window. Rebuild Server 1 as standby, validate replication, decide whether Server 2 remains primary, and avoid unnecessary role reversal.")

add_heading(doc, "15. Production validation and acceptance", 1)
add_heading(doc, "15.1 Service validation", 2)
add_code(doc, """
sudo systemctl is-active postgresql
sudo systemctl is-active prometheus
sudo systemctl is-active alertmanager
sudo systemctl is-active node_exporter
sudo systemctl is-active haproxy

curl -fsS http://localhost:9090/-/ready
curl -fsS http://localhost:9093/-/ready
curl -fsS http://localhost:9100/metrics | head
curl -fsS http://127.0.0.1:9091/-/ready
""")

add_heading(doc, "15.2 Acceptance checklist", 2)
checks = [
    "Both Prometheus replicas show all required targets as UP.",
    "Both Prometheus replicas evaluate the same alert-rule set.",
    "Both Alertmanager nodes show the expected two-member cluster.",
    "A test alert is delivered and acknowledged through the approved notification channel.",
    "Grafana reports database status `ok` and loads dashboards through HAProxy.",
    "Stopping Prometheus A causes Grafana queries to continue through Prometheus B.",
    "Stopping Alertmanager A does not prevent notification delivery.",
    "PostgreSQL standby reports `pg_is_in_recovery() = true` and acceptable replay lag.",
    "A controlled PostgreSQL promotion test has been completed in a safe test window.",
    "Grafana DNS and TLS certificate validation succeed.",
    "Default administrator credentials have been replaced.",
    "Anonymous access and user self-registration are disabled.",
    "All exporter and metrics ports are restricted to approved monitoring sources.",
    "Backups complete successfully and a restoration test has passed.",
    "The operations team has the runbook, ownership matrix, and escalation contacts.",
]
for check in checks:
    add_bullet(doc, "☐ " + check)

add_heading(doc, "15.3 Failure tests", 2)
add_table(
    doc,
    ["Test", "Expected result", "Evidence"],
    [
        ("Stop local Prometheus", "HAProxy uses remote replica; dashboards continue.", "HAProxy status/log and Grafana query"),
        ("Stop one Alertmanager", "Surviving member sends test notification.", "Notification and cluster status"),
        ("Block one Node Exporter", "TargetDown fires after configured delay.", "Prometheus alert and notification"),
        ("Restart monitoring node", "Enabled services return automatically.", "Systemd status and logs"),
        ("PostgreSQL standby lag", "Alert before RPO is exceeded.", "Replication dashboard and alert"),
        ("Restore PostgreSQL backup", "Grafana starts with expected data.", "Restore record and screenshots"),
    ],
    [2300, 4500, 2560],
)

page_break(doc)

add_heading(doc, "16. Operations, patching, and governance", 1)
add_heading(doc, "16.1 Daily checks", 2)
for item in [
    "Target availability and scrape errors.",
    "Prometheus disk usage, ingestion rate, query latency, and rule failures.",
    "Alertmanager cluster membership and notification errors.",
    "PostgreSQL replication state and replay lag.",
    "Grafana health, failed logins, data-source errors, and certificate expiry.",
]:
    add_bullet(doc, item)

add_heading(doc, "16.2 Upgrade sequence", 2)
for number, text in enumerate([
    "Review release notes, security advisories, compatibility, and rollback requirements.",
    "Back up PostgreSQL and all configuration files.",
    "Upgrade and validate Server 2 monitoring components first.",
    "Confirm Prometheus and Alertmanager redundancy remains healthy.",
    "Upgrade Server 1 during the approved window.",
    "For Grafana, upgrade the standby package first but keep it stopped.",
    "Upgrade active Grafana only after a verified database backup.",
    "Do not downgrade Grafana without restoring the pre-upgrade database backup.",
], start=1):
    add_number(doc, text, number)

add_heading(doc, "16.3 Configuration change sequence", 2)
add_code(doc, """
sudo promtool check config /etc/prometheus/prometheus.yml
sudo promtool check rules /etc/prometheus/rules/*.yml
sudo amtool check-config /etc/alertmanager/alertmanager.yml
sudo haproxy -c -f /etc/haproxy/haproxy.cfg
sudo nginx -t
""")
add_body(doc, "Deploy to Server 2 first, validate, then deploy the identical functional change to Server 1. Preserve only the expected node-specific values, such as replica labels and peer addresses.")

add_heading(doc, "16.4 Ownership", 2)
add_table(
    doc,
    ["Area", "Primary owner", "Required responsibilities"],
    [
        ("RHEL and storage", "Linux/platform team", "Patching, capacity, filesystem, service recovery"),
        ("PostgreSQL", "Database/platform team", "Replication, backup, promotion, recovery"),
        ("Grafana", "Monitoring team", "Users, dashboards, provisioning, upgrades"),
        ("Prometheus/Alertmanager", "Monitoring team", "Scrapes, rules, retention, notifications"),
        ("Application endpoints", "Application owners", "Metrics enablement and version compatibility"),
        ("Network/PKI/DNS", "Infrastructure teams", "Firewall, certificates, names, failover"),
        ("Incident response", "Operations", "Acknowledgment, escalation, runbook execution"),
    ],
    [1950, 2200, 5210],
)

page_break(doc)

add_heading(doc, "Appendix A — Command and configuration inventory", 1)
add_table(
    doc,
    ["Component", "Configuration", "Data", "Service"],
    [
        ("PostgreSQL", "/var/lib/pgsql/data/*.conf", "/var/lib/pgsql/data", "postgresql"),
        ("Grafana", "/etc/grafana/grafana.ini", "PostgreSQL database", "grafana-server"),
        ("Grafana provisioning", "/etc/grafana/provisioning", "N/A", "grafana-server"),
        ("Prometheus", "/etc/prometheus", "/var/lib/prometheus", "prometheus"),
        ("Alertmanager", "/etc/alertmanager", "/var/lib/alertmanager", "alertmanager"),
        ("Node Exporter", "systemd unit", "N/A", "node_exporter"),
        ("HAProxy", "/etc/haproxy/haproxy.cfg", "N/A", "haproxy"),
        ("NGINX", "/etc/nginx/conf.d/grafana.conf", "TLS files", "nginx"),
    ],
    [2100, 2900, 2100, 2260],
)

add_heading(doc, "Appendix B — Troubleshooting quick reference", 1)
add_table(
    doc,
    ["Symptom", "Checks"],
    [
        ("Grafana database error", "Test PostgreSQL TLS/login; inspect Grafana logs; verify CA, DNS, password file, pg_hba.conf."),
        ("Grafana dashboard has no data", "Test HAProxy 9091; check both Prometheus readiness endpoints; inspect target status."),
        ("Prometheus will not start", "Run promtool; check ownership, retention value, disk space, and systemd logs."),
        ("Target DOWN", "Test endpoint from both Prometheus nodes; check firewall, authentication, certificate, and application health."),
        ("Duplicate notifications", "Inspect Alertmanager cluster membership and network partition; duplicates may be intentional fail-open behavior."),
        ("PostgreSQL standby not streaming", "Check `.pgpass`, replication slot, pg_hba.conf, certificate, port 5432, and primary logs."),
        ("NGINX 502", "Confirm Grafana is listening on 127.0.0.1:3000 and SELinux boolean permits proxy connection."),
        ("Standby Grafana cannot decrypt secrets", "Confirm the same Grafana secret_key file is deployed to both nodes."),
    ],
    [2800, 6560],
)

add_heading(doc, "Appendix C — Authoritative references", 1)
refs = [
    ("Grafana RHEL installation", "https://grafana.com/docs/grafana/latest/setup-grafana/installation/redhat-rhel-fedora/"),
    ("Grafana configuration", "https://grafana.com/docs/grafana/latest/setup-grafana/configure-grafana/"),
    ("Grafana high availability", "https://grafana.com/docs/grafana/latest/setup-grafana/set-up-for-high-availability/"),
    ("Grafana backup", "https://grafana.com/docs/grafana/latest/administration/back-up-grafana/"),
    ("Prometheus downloads", "https://prometheus.io/download/"),
    ("Prometheus storage", "https://prometheus.io/docs/prometheus/latest/storage/"),
    ("Prometheus security model", "https://prometheus.io/docs/operating/security/"),
    ("Alertmanager high availability", "https://prometheus.io/docs/alerting/latest/high_availability/"),
    ("RHEL 9 PostgreSQL installation", "https://docs.redhat.com/en/documentation/red_hat_enterprise_linux/9/html/configuring_and_using_database_servers/using-postgresql_configuring-and-using-database-servers"),
    ("PostgreSQL streaming replication", "https://www.postgresql.org/docs/16/warm-standby.html"),
    ("PostgreSQL base backup", "https://www.postgresql.org/docs/current/app-pgbasebackup.html"),
    ("GitLab Prometheus monitoring", "https://docs.gitlab.com/administration/monitoring/prometheus/"),
    ("Jenkins Prometheus plugin", "https://plugins.jenkins.io/prometheus/"),
    ("SonarQube Web API", "https://docs.sonarsource.com/sonarqube-server/extension-guide/web-api"),
    ("Nexus Prometheus metrics", "https://help.sonatype.com/en/prometheus.html"),
]
for label, url in refs:
    p = doc.add_paragraph(style="Normal")
    r = p.add_run(label + ": ")
    set_run_font(r, bold=True, size=10)
    r = p.add_run(url)
    set_run_font(r, size=9.5, color=BLUE2)

add_callout(doc, "Version note", "This runbook selected Prometheus 3.5.5 LTS, Alertmanager 0.33.1, and Node Exporter 1.12.1 based on the official Prometheus download catalog available on 23 July 2026. Revalidate versions and checksums immediately before production installation.", "note")

# Document properties and final XML settings.
doc.core_properties.title = "Grafana Monitoring Platform Production Installation Runbook"
doc.core_properties.subject = "Two-server RHEL 9 production deployment"
doc.core_properties.author = "Platform Engineering"
doc.core_properties.keywords = "Grafana, Prometheus, Alertmanager, PostgreSQL, RHEL 9"

settings = doc.settings._element
update_fields = settings.find(qn("w:updateFields"))
if update_fields is None:
    update_fields = OxmlElement("w:updateFields")
    settings.append(update_fields)
update_fields.set(qn("w:val"), "true")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
