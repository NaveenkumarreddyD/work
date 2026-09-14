from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLUE = "244B65"
LIGHT_BLUE = "EAF1F5"
GRAY = "5F6368"
BLACK = "111111"
WHITE = "FFFFFF"


def setup_document(title, subtitle=None):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, before, after in (
        ("Heading 1", 14, 12, 4),
        ("Heading 2", 11.5, 9, 3),
    ):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLUE)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_after = Pt(3)
    title_p.paragraph_format.keep_with_next = True
    title_run = title_p.add_run(title)
    title_run.font.name = "Arial"
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor.from_string(BLUE)

    if subtitle:
        subtitle_p = doc.add_paragraph()
        subtitle_p.paragraph_format.space_after = Pt(12)
        subtitle_run = subtitle_p.add_run(subtitle)
        subtitle_run.font.name = "Arial"
        subtitle_run.font.size = Pt(10.5)
        subtitle_run.font.color.rgb = RGBColor.from_string(GRAY)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Maximo Manage attachment storage work")
    footer_run.font.name = "Arial"
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor.from_string(GRAY)

    add_numbering(doc, 30, bullet=True)
    add_numbering(doc, 31, bullet=False)
    add_numbering(doc, 32, bullet=False)
    return doc


def add_numbering(doc, num_id, *, bullet=False):
    numbering = doc.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(num_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), "bullet" if bullet else "decimal")
    level.append(fmt)
    text = OxmlElement("w:lvlText")
    text.set(qn("w:val"), "•" if bullet else "%1.")
    level.append(text)
    p_pr = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    level.append(p_pr)
    abstract.append(level)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_id = OxmlElement("w:abstractNumId")
    abstract_id.set(qn("w:val"), str(num_id))
    num.append(abstract_id)
    numbering.append(num)


def add_list(doc, items, *, numbered=False, num_id=None):
    selected_num_id = num_id if num_id is not None else (31 if numbered else 30)
    for item in items:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.38)
        paragraph.paragraph_format.first_line_indent = Inches(-0.19)
        paragraph.paragraph_format.space_after = Pt(4)
        p_pr = paragraph._p.get_or_add_pPr()
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num_id_node = OxmlElement("w:numId")
        num_id_node.set(qn("w:val"), str(selected_num_id))
        num_pr.append(ilvl)
        num_pr.append(num_id_node)
        p_pr.append(num_pr)
        paragraph.add_run(item)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for side, value in (("top", 100), ("start", 110), ("bottom", 100), ("end", 110)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tc_pr.append(margins)


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    indent = OxmlElement("w:tblInd")
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")
    tbl_pr.append(indent)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = OxmlElement("w:tcW")
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            tc_pr.append(tc_w)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_text(cell, text, *, bold=False, color=BLACK, size=9.2):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_property_table(doc):
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    for index, text in enumerate(("Property", "Value used")):
        set_cell_shading(table.rows[0].cells[index], BLUE)
        set_cell_text(table.rows[0].cells[index], text, bold=True, color=WHITE)
    properties = [
        ("mxe.cosaccesskey", "Configured in Manage; value not included here"),
        ("mxe.cossecretkey", "Configured in Manage; value not included here"),
        ("mxe.cosendpointuri", "https://bhm-pwrsclnfs.lac1.biz:9021"),
        ("mxe.cosbucketname", "dr-maximo-bckt"),
        ("mxe.attachmentstorage", "com.ibm.tivoli.maximo.oslc.provider.COSAttachmentStorage"),
        ("mxe.doclink.securedAttachment", "true"),
        ("mxe.doclink.doctypes.defpath", "cos:doclinks/default"),
        ("mxe.doclink.doctypes.topLevelPaths", "cos:doclinks"),
        (
            "mxe.doclink.path01",
            "cos:doclinks=https://drgitopswks.manage.drgitopsapp.apps.drroc4.lac1.biz/maximo/oslc/cosdoclink",
        ),
    ]
    for prop, value in properties:
        cells = table.add_row().cells
        set_cell_text(cells[0], prop)
        set_cell_text(cells[1], value)
    set_table_geometry(table, [3500, 6100])


def add_reference(doc, label, url):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(f"{label}: {url}")
    run.font.name = "Arial"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(GRAY)


def add_code_block(doc, code):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [9600])
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F4F4F4")
    cell.text = ""
    for line in code.strip("\n").splitlines():
        paragraph = cell.add_paragraph() if cell.paragraphs[0].text else cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        run = paragraph.add_run(line or " ")
        run.font.name = "Courier New"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


# Document 1: configuration
config_doc = setup_document(
    "Manage S3 Attachment Configuration",
    "Summary of the PowerScale S3 setup and the test completed in Manage",
)
config_doc.add_paragraph(
    "Manage was using the /doclinks file system on NFS for attachments. We configured the environment to use the PowerScale S3-compatible bucket for new attachments."
)

config_doc.add_heading("Environment", level=1)
add_list(
    config_doc,
    [
        "Manage version: 8.7.24",
        "Instance and workspace: drgitopsapp / drgitopswks",
        "S3 endpoint: https://bhm-pwrsclnfs.lac1.biz:9021",
        "S3 bucket: dr-maximo-bckt",
    ],
)

config_doc.add_heading("What we changed", level=1)
add_list(
    config_doc,
    [
        "Added the PowerScale root and subordinate CA certificates so the Manage pods trust the HTTPS endpoint.",
        "Configured the bucket, endpoint, access key, and secret key used by Manage.",
        "Changed the attachment provider from file storage to the Maximo S3 attachment provider.",
        "Updated the doclink paths from /doclinks to cos:doclinks.",
        "Restarted the Manage server bundle pods so the new values were loaded.",
    ],
)

config_doc.add_heading("Manage properties", level=1)
add_property_table(config_doc)

config_doc.add_page_break()
config_doc.add_heading("Import the PowerScale S3 certificates", level=1)
config_doc.add_paragraph(
    "Manage must trust the certificate chain used by the PowerScale HTTPS endpoint. Store the Root CA and SubCA in the configured secret manager and let the GitOps rendering process add them to the Manage truststore. The endpoint supplies the leaf certificate, so only the issuing Root CA and SubCA are imported."
)

config_doc.add_heading("1. Check the certificate files", level=2)
add_code_block(
    config_doc,
    """
openssl x509 -in RootCA.cer -noout -subject -issuer -dates
openssl x509 -in SubCA.cer  -noout -subject -issuer -dates
""",
)
config_doc.add_paragraph(
    "The files must be PEM certificates and contain BEGIN CERTIFICATE and END CERTIFICATE lines. Rename them to powerscale-s3-rootca.pem and powerscale-s3-subca.pem after checking them."
)

config_doc.add_heading("2. Store the certificates in the secret manager", level=2)
config_doc.add_paragraph(
    "Update the environment's existing S3 secret record through the approved secret-management process. Keep the field names stable so the same GitOps template works when the secret backend changes. Store the complete PEM text, including the BEGIN CERTIFICATE and END CERTIFICATE lines."
)
add_list(
    config_doc,
    [
        "powerscale_s3_subca: complete subordinate CA certificate in PEM format.",
        "powerscale_s3_rootca: complete root CA certificate in PEM format.",
        "Keep the S3 endpoint, bucket, access key, and secret key in the same protected environment record or in separately controlled secret records.",
        "Do not place certificate contents, access keys, secret keys, or secret-manager tokens in Git.",
    ],
)

config_doc.add_heading("3. Reference the certificates in Manage", level=2)
config_doc.add_paragraph(
    "Map the two secret fields into the importedCerts section of the rendered ManageWorkspace configuration. The placeholders below represent values supplied by the configured secret integration; they must resolve to the PEM text before the resource is applied."
)
add_code_block(
    config_doc,
    """
settings:
  deployment:
    importedCerts:
      - alias: powerscale-s3-subca
        crt: |
          ${POWERSCALE_S3_SUBCA_PEM}
      - alias: powerscale-s3-rootca
        crt: |
          ${POWERSCALE_S3_ROOTCA_PEM}
""",
)
config_doc.add_paragraph(
    "The current secret integration can resolve these values from its own backend. When the backend changes, update only the secret mapping or renderer, not the ManageWorkspace structure. Use MANAGE_ATTACHMENT_PROVIDER=s3-migration while the NFS migration and rollback window are open. Render and commit the configuration, then allow Argo CD to synchronize it."
)

config_doc.add_heading("4. Verify the imported certificate aliases", level=2)
add_code_block(
    config_doc,
    """
export INSTANCE_ID=drgitopsapp
export WORKSPACE_ID=drgitopswks
export MANAGE_NS=mas-${INSTANCE_ID}-manage
export MANAGEWORKSPACE=${INSTANCE_ID}-${WORKSPACE_ID}

oc get manageworkspace "$MANAGEWORKSPACE" -n "$MANAGE_NS" \\
  -o jsonpath='{.spec.settings.deployment.importedCerts[*].alias}{"\\n"}'
""",
)
config_doc.add_paragraph(
    "The output should contain powerscale-s3-subca and powerscale-s3-rootca. A successful attachment upload to the HTTPS S3 endpoint also confirms that the running Manage workload trusts the chain."
)

config_doc.add_heading("Test result", level=1)
config_doc.add_paragraph(
    "After the restart, we uploaded a new attachment from Manage. The attachment was created in the PowerScale S3 bucket instead of the NFS /doclinks location. This confirms that new attachment writes are using S3."
)
add_list(
    config_doc,
    [
        "The attachment upload completed in Manage.",
        "The new object was visible in dr-maximo-bckt.",
        "The database reference uses the cos:doclinks path.",
    ],
)

config_doc.add_heading("Notes", level=1)
add_list(
    config_doc,
    [
        "Access keys, secret keys, secret-manager tokens, and private keys are not included in this document.",
        "Use the Manage System Properties application for routine property changes. Direct database updates should be limited to controlled testing or recovery.",
        "mxe.cosnestedfile is not needed for this sample because the migrated object names were stored at the bucket root.",
    ],
)

config_doc.add_heading("Two ways to update the properties", level=1)
config_doc.add_paragraph(
    "Both methods update the same global values. The Manage UI is easier for normal administration. DBeaver is useful for a controlled test or recovery when the existing rows are already present."
)

config_doc.add_heading("Option 1 - Manage System Properties", level=2)
add_list(
    config_doc,
    [
        "Open System Configuration > Platform Configuration > System Properties.",
        "Search for mxe.cos, mxe.attachmentstorage, and mxe.doclink.",
        "Open each property and enter the Global Value shown in the table above.",
        "Save each property. Do not record the access key or secret key in notes or screenshots.",
        "Restart all Manage server bundle pods and confirm the Current Value is updated.",
    ],
    numbered=True,
)

config_doc.add_heading("Option 2 - DBeaver", level=2)
config_doc.add_paragraph(
    "Turn Auto-commit off before starting. These commands update existing COMMON rows only. They do not create missing properties and they do not encrypt the access key or secret key. DBeaver prompts for the two credential parameters."
)

config_doc.add_heading("Check the property flags and current values", level=2)
add_code_block(
    config_doc,
    """
SELECT propname, encrypted, masked
FROM maximo.maxprop
WHERE LOWER(propname) IN ('mxe.cosaccesskey', 'mxe.cossecretkey');
""",
)

config_doc.add_heading("Read the current values", level=2)
add_code_block(
    config_doc,
    """
SELECT propname,
       servername,
       CASE
           WHEN LOWER(propname) IN ('mxe.cosaccesskey', 'mxe.cossecretkey')
           THEN '<configured>'
           ELSE propvalue
       END AS propvalue
FROM maximo.maxpropvalue
WHERE servername = 'COMMON'
  AND LOWER(propname) IN (
      'mxe.cosaccesskey',
      'mxe.cossecretkey',
      'mxe.cosendpointuri',
      'mxe.cosbucketname',
      'mxe.attachmentstorage',
      'mxe.doclink.securedattachment',
      'mxe.doclink.doctypes.defpath',
      'mxe.doclink.doctypes.toplevelpaths',
      'mxe.doclink.path01'
  )
ORDER BY propname;
""",
)

config_doc.add_heading("Update the existing rows", level=2)
add_code_block(
    config_doc,
    """
SAVEPOINT before_s3_properties;

UPDATE maximo.maxpropvalue SET propvalue = :s3_access_key
 WHERE servername = 'COMMON' AND LOWER(propname) = 'mxe.cosaccesskey';

UPDATE maximo.maxpropvalue SET propvalue = :s3_secret_key
 WHERE servername = 'COMMON' AND LOWER(propname) = 'mxe.cossecretkey';

UPDATE maximo.maxpropvalue SET propvalue = 'https://bhm-pwrsclnfs.lac1.biz:9021'
 WHERE servername = 'COMMON' AND LOWER(propname) = 'mxe.cosendpointuri';

UPDATE maximo.maxpropvalue SET propvalue = 'dr-maximo-bckt'
 WHERE servername = 'COMMON' AND LOWER(propname) = 'mxe.cosbucketname';

UPDATE maximo.maxpropvalue
 SET propvalue = 'com.ibm.tivoli.maximo.oslc.provider.COSAttachmentStorage'
 WHERE servername = 'COMMON' AND LOWER(propname) = 'mxe.attachmentstorage';

UPDATE maximo.maxpropvalue SET propvalue = 'true'
 WHERE servername = 'COMMON'
   AND LOWER(propname) = 'mxe.doclink.securedattachment';

UPDATE maximo.maxpropvalue SET propvalue = 'cos:doclinks/default'
 WHERE servername = 'COMMON'
   AND LOWER(propname) = 'mxe.doclink.doctypes.defpath';

UPDATE maximo.maxpropvalue SET propvalue = 'cos:doclinks'
 WHERE servername = 'COMMON'
   AND LOWER(propname) = 'mxe.doclink.doctypes.toplevelpaths';

UPDATE maximo.maxpropvalue
 SET propvalue = 'cos:doclinks=https://drgitopswks.manage.drgitopsapp.apps.drroc4.lac1.biz/maximo/oslc/cosdoclink'
 WHERE servername = 'COMMON' AND LOWER(propname) = 'mxe.doclink.path01';
""",
)

config_doc.add_heading("Verify and commit", level=2)
config_doc.add_paragraph(
    "DBeaver should report one updated row for each statement. Run the masked SELECT above again and check all nine values before committing. After committing, restart every Manage server bundle. If a property update reports zero rows, roll back and create or update that property through the Manage UI instead of inserting directly into MAXPROP or MAXPROPVALUE."
)
add_code_block(
    config_doc,
    """
COMMIT;

-- Use this instead of COMMIT if the values are wrong:
-- ROLLBACK TO before_s3_properties;
    """,
)
config_doc.save("Manage_S3_Attachment_Configuration.docx")


# Document 2: migration and rollback
migration_doc = setup_document(
    "NFS to S3 Attachment Migration Test",
    "Controlled test of existing file-based attachments and the rollback steps",
)
migration_doc.add_paragraph(
    "We used a small set of attachments to test IBM's file-to-S3 migration process before considering a larger migration. The source files were on the NFS /doclinks mount and the target was the PowerScale S3 bucket."
)

migration_doc.add_heading("Test sample", level=1)
add_list(
    migration_doc,
    [
        "33 attachments linked to Work Orders, Assets, and Locations.",
        "File types included PDF, JPG, PNG, TXT, CSV, and XML.",
        "DOCINFOID 3676587 and 3676590 through 3676621 were included.",
        "The original paths were under /doclinks/attachments and /doclinks/diagrams.",
    ],
)

migration_doc.add_heading("Detailed steps used", level=1)

migration_doc.add_heading("1. Create the test sample", level=2)
migration_doc.add_paragraph(
    "We uploaded files through Manage while attachments were still using NFS. The files were spread across Work Orders, Assets, and a Location instead of using one record. We confirmed they opened before starting the migration."
)
add_list(
    migration_doc,
    [
        "Work Order owner IDs used: 14179087, 14188816, 14188818, 14188819, and 14188820.",
        "Asset owner IDs used: 14260000 and 14260319.",
        "Location owner ID used: 1851554.",
    ],
)

migration_doc.add_heading("2. Save the before-migration database results", level=2)
migration_doc.add_paragraph(
    "Run this query in DBeaver and export the result to CSV. Keep this file because it contains the original /doclinks paths needed for rollback."
)
add_code_block(
    migration_doc,
    """
SELECT DISTINCT
       di.docinfoid,
       di.document,
       di.urlname,
       di.createdate,
       dl.ownertable,
       dl.ownerid
FROM maximo.docinfo di
JOIN maximo.doclinks dl ON dl.docinfoid = di.docinfoid
WHERE di.docinfoid IN (
  3676587,3676590,3676591,3676592,3676593,3676594,3676595,
  3676596,3676597,3676598,3676599,3676600,3676601,3676602,
  3676603,3676604,3676605,3676606,3676607,3676608,3676609,
  3676610,3676611,3676612,3676613,3676614,3676615,3676616,
  3676617,3676618,3676619,3676620,3676621
)
ORDER BY di.docinfoid, dl.ownertable, dl.ownerid;
""",
)

migration_doc.add_heading("3. Open the Manage pod that contains the migration tool", level=2)
add_code_block(
    migration_doc,
    """
export INSTANCE_ID=drgitopsapp
export WORKSPACE_ID=drgitopswks
export MANAGE_NS=mas-${INSTANCE_ID}-manage

oc get pods -n "$MANAGE_NS"

# Choose the Manage admin/tools pod that has file2s3.sh and /doclinks mounted.
export MANAGE_POD=<manage-tools-pod>
oc rsh -n "$MANAGE_NS" "$MANAGE_POD"

FILE2S3=$(find / -type f -name file2s3.sh 2>/dev/null | head -1)
echo "$FILE2S3"
ls -ld /doclinks /doclinks/attachments /doclinks/diagrams
""",
)
migration_doc.add_paragraph(
    "Do not continue if file2s3.sh is not found or the pod cannot see the original /doclinks files."
)

migration_doc.add_heading("4. Set the S3 values in the pod", level=2)
migration_doc.add_paragraph(
    "Use read prompts for the credentials so they are not written into the shell command or this document."
)
add_code_block(
    migration_doc,
    """
export S3_ENDPOINT=https://bhm-pwrsclnfs.lac1.biz:9021
export S3_BUCKET=dr-maximo-bckt

read -r -p "S3 access key: " S3_ACCESS_KEY
read -r -s -p "S3 secret key: " S3_SECRET_KEY
echo
export S3_ACCESS_KEY S3_SECRET_KEY
""",
)

migration_doc.add_heading("5. Run the controlled migration", level=2)
add_code_block(
    migration_doc,
    """
DOCINFO_IDS="3676587,3676590,3676591,3676592,3676593,3676594"
DOCINFO_IDS="${DOCINFO_IDS},3676595,3676596,3676597,3676598,3676599"
DOCINFO_IDS="${DOCINFO_IDS},3676600,3676601,3676602,3676603,3676604"
DOCINFO_IDS="${DOCINFO_IDS},3676605,3676606,3676607,3676608,3676609"
DOCINFO_IDS="${DOCINFO_IDS},3676610,3676611,3676612,3676613,3676614"
DOCINFO_IDS="${DOCINFO_IDS},3676615,3676616,3676617,3676618,3676619"
DOCINFO_IDS="${DOCINFO_IDS},3676620,3676621"

bash "$FILE2S3" \\
  -b"$S3_BUCKET" \\
  -x"$S3_ACCESS_KEY" \\
  -s"$S3_SECRET_KEY" \\
  -r"$S3_ENDPOINT" \\
  -w"DOCINFOID/**/IN(${DOCINFO_IDS})"

unset S3_ACCESS_KEY S3_SECRET_KEY
""",
)
migration_doc.add_paragraph(
    "Keep the filter exactly in the DOCINFOID/**/IN(...) format without spaces around IN. A filter written as IN (...) caused the tool to generate invalid Oracle SQL during the first attempt."
)

migration_doc.add_heading("6. Check the migration output", level=2)
add_list(
    migration_doc,
    [
        "Each selected DOCINFOID should show that the file was uploaded.",
        "Each record should show an update from /doclinks/... to cos:doclinks/....",
        "The final message should say FileToS3Move completed without errors.",
        "Stop and investigate if the output contains ORA-, SSL, certificate, 403, bucket, or upload errors.",
    ],
)

migration_doc.add_heading("7. Save the after-migration database results", level=2)
migration_doc.add_paragraph(
    "Run the same DBeaver query from step 2 and export it again with the same column headers. The URLNAME values should now start with cos:doclinks. Keep both exports together."
)

migration_doc.add_heading("8. Apply the S3 settings and restart Manage", level=2)
migration_doc.add_paragraph(
    "Configure the nine Manage properties and the PowerScale certificate chain using the steps in Manage_S3_Attachment_Configuration.docx. The properties can be set through the Manage UI or with the DBeaver commands in that document."
)
add_code_block(
    migration_doc,
    """
export MANAGE_NS=mas-drgitopsapp-manage

oc get pods -n "$MANAGE_NS" -o name \\
  | grep -E 'drgitopswks-(ui|cron|mea|report)-'

oc get pods -n "$MANAGE_NS" -o name \\
  | grep -E 'drgitopswks-(ui|cron|mea|report)-' \\
  | xargs oc delete -n "$MANAGE_NS"

oc get pods -n "$MANAGE_NS" -w
""",
)

migration_doc.add_heading("9. Run the post-migration test", level=2)
add_list(
    migration_doc,
    [
        "Open a migrated record in Manage and confirm its attachment can be opened or downloaded.",
        "Upload a new attachment from Manage.",
        "Confirm the new DOCINFO.URLNAME starts with cos:doclinks.",
        "Confirm the new object is visible in dr-maximo-bckt.",
        "Review the Manage logs for TLS, authentication, bucket, or object errors.",
    ],
    numbered=True,
)
add_code_block(
    migration_doc,
    """
SELECT docinfoid, document, urlname, createdate
FROM maximo.docinfo
WHERE LOWER(urlname) LIKE 'cos:doclinks%'
ORDER BY createdate DESC
FETCH FIRST 20 ROWS ONLY;
""",
)

migration_doc.add_heading("Result", level=1)
migration_doc.add_paragraph(
    "The controlled migration completed successfully. The selected database records now point to S3. After the S3 settings were applied, we also uploaded a new attachment from Manage and confirmed that the new object was stored in the PowerScale S3 bucket."
)
add_list(
    migration_doc,
    [
        "33 selected attachments were processed by the migration tool.",
        "The migration finished without errors.",
        "The selected URLNAME values now use cos:doclinks.",
        "A new post-migration attachment was stored in S3.",
    ],
)

migration_doc.add_page_break()
migration_doc.add_heading("Rollback test", level=1)
migration_doc.add_paragraph(
    "The rollback test should use only a few of the migrated records. Keep both the NFS files and S3 objects until the test is finished."
)
rollback_steps = [
    "Export the current S3 property values and keep the previous NFS property values.",
    "Keep the before-and-after URLNAME mapping for the selected DOCINFOID records.",
    "Confirm the original files still exist on NFS and pause new attachment uploads during the rollback test.",
    "Choose three to five records across Work Orders, Assets, and Locations.",
    "Restore the previous NFS attachment values through the Manage System Properties application or by updating the existing COMMON rows in MAXPROPVALUE with the saved values.",
    "Restore the selected DOCINFO.URLNAME values to their original /doclinks paths using the saved mapping.",
    "Restart all Manage server bundle pods.",
    "Open and download the selected attachments, then upload one new file and confirm it is written to NFS.",
    "Reapply the S3 properties through the UI or DBeaver, restore the selected cos:doclinks URLNAME values, restart the pods, and test S3 again.",
]
add_list(migration_doc, rollback_steps, numbered=True, num_id=32)

migration_doc.add_heading("Rollback note", level=2)
migration_doc.add_paragraph(
    "A full rollback after users start adding files to S3 needs an extra step. Any file that exists only in S3 must be copied back to NFS, and its DOCINFO.URLNAME must be changed to the correct NFS path. Changing the system properties alone is not enough."
)

migration_doc.add_heading("Rollback is successful when", level=2)
add_list(
    migration_doc,
    [
        "The selected attachments open from NFS.",
        "A new attachment can be written to NFS.",
        "No DOCINFO record points to a missing file.",
        "The same records work again after the environment is returned to S3.",
    ],
)

migration_doc.save("NFS_to_S3_Attachment_Migration_Test.docx")

print("Manage_S3_Attachment_Configuration.docx")
print("NFS_to_S3_Attachment_Migration_Test.docx")
