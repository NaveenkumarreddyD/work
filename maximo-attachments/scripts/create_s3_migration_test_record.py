from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = "B-47623_Maximo_NFS_to_S3_Migration_Test_Record.docx"

NAVY = "18324A"
TEAL = "0E5A64"
GOLD = "C89B3C"
PALE_TEAL = "E8F3F4"
PALE_GOLD = "F8F1E2"
LIGHT_GRAY = "F3F5F6"
MID_GRAY = "6B747C"
WHITE = "FFFFFF"


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_text(cell, text, *, bold=False, color=NAVY, size=9, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MID_GRAY)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    relation_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), TEAL)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(color)
    run_properties.append(underline)
    run.append(run_properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_numbering_definition(document, num_id, *, bullet=False):
    numbering = document.part.numbering_part.element
    abstract_id = num_id
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if bullet else "decimal")
    level.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if bullet else "%1.")
    level.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    level.append(lvl_jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    level.append(p_pr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num_id = OxmlElement("w:abstractNumId")
    abstract_num_id.set(qn("w:val"), str(abstract_id))
    num.append(abstract_num_id)
    numbering.append(num)


def add_list_item(document, text, num_id, *, bold_lead=None):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.38)
    paragraph.paragraph_format.first_line_indent = Inches(-0.19)
    paragraph.paragraph_format.space_after = Pt(3)
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_node)
    p_pr.append(num_pr)
    if bold_lead and text.startswith(bold_lead):
        first = paragraph.add_run(bold_lead)
        first.bold = True
        paragraph.add_run(text[len(bold_lead):])
    else:
        paragraph.add_run(text)
    return paragraph


def add_heading(document, text, level=1):
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.78)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.32)
section.footer_distance = Inches(0.35)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(9.5)
normal.font.color.rgb = RGBColor.from_string(NAVY)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.08

for name, size, color, before, after in (
    ("Heading 1", 14, NAVY, 12, 4),
    ("Heading 2", 11, TEAL, 8, 3),
    ("Heading 3", 9.5, NAVY, 6, 2),
):
    style = styles[name]
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

header = section.header.paragraphs[0]
header.text = "IBM MAXIMO MANAGE  |  ATTACHMENT STORAGE TEST"
header.alignment = WD_ALIGN_PARAGRAPH.LEFT
header.paragraph_format.space_after = Pt(0)
for run in header.runs:
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(TEAL)

footer = section.footer.paragraphs[0]
left = footer.add_run("Internal test record  |  Jira B-47623")
left.font.name = "Arial"
left.font.size = Pt(8)
left.font.color.rgb = RGBColor.from_string(MID_GRAY)
footer.add_run("\t")
add_page_number(footer)

add_numbering_definition(doc, 20, bullet=True)
add_numbering_definition(doc, 21, bullet=False)
add_numbering_definition(doc, 22, bullet=False)

title = doc.add_paragraph()
title.paragraph_format.space_after = Pt(3)
title.paragraph_format.keep_with_next = True
run = title.add_run("Maximo Manage NFS to S3 Attachment Migration")
run.font.name = "Arial"
run.font.size = Pt(22)
run.font.bold = True
run.font.color.rgb = RGBColor.from_string(NAVY)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(12)
subtitle.paragraph_format.keep_with_next = True
run = subtitle.add_run("Test record, validation checklist, and rollback plan")
run.font.name = "Arial"
run.font.size = Pt(11)
run.font.color.rgb = RGBColor.from_string(TEAL)

summary = doc.add_table(rows=5, cols=2)
summary.style = "Table Grid"
set_table_geometry(summary, [2100, 7260])
for row_index, (label, value) in enumerate(
    [
        ("Jira", "B-47623 - Configure Manage S3 attachment storage"),
        ("Status", "In progress - migration completed; final restart validation and rollback test pending"),
        ("Environment", "drgitopsapp / drgitopswks | Maximo Manage 8.7.24 | Oracle"),
        ("Storage", "PowerScale NFS /doclinks to PowerScale S3-compatible storage"),
        ("Prepared", "August 5, 2026"),
    ]
):
    cells = summary.rows[row_index].cells
    set_cell_shading(cells[0], LIGHT_GRAY)
    set_cell_text(cells[0], label, bold=True, color=TEAL)
    set_cell_text(cells[1], value)

doc.add_paragraph()
callout = doc.add_table(rows=1, cols=1)
callout.style = "Table Grid"
set_table_geometry(callout, [9360])
set_cell_shading(callout.cell(0, 0), PALE_TEAL)
set_cell_text(
    callout.cell(0, 0),
    "Current result: 33 controlled test attachments were copied to S3 and their DOCINFO.URLNAME values were updated to cos:doclinks. The IBM migration tool completed without errors.",
    bold=True,
    color=TEAL,
    size=9.5,
)

add_heading(doc, "1. Objective")
doc.add_paragraph(
    "Confirm that existing Maximo Manage attachments can be migrated from file-based NFS storage to PowerScale S3, remain accessible from Manage, and be recovered to NFS if the S3 configuration fails. This is a controlled non-production test and does not include a production cutover."
)

add_heading(doc, "2. Test Scope")
for item in [
    "33 test attachments linked to Work Orders, Assets, and Locations.",
    "File types included PDF, JPG, PNG, TXT, CSV, and XML.",
    "Selected records: DOCINFOID 3676587 and 3676590 through 3676621.",
    "Source paths were under /doclinks/attachments and /doclinks/diagrams.",
    "Target bucket: dr-maximo-bckt at https://bhm-pwrsclnfs.lac1.biz:9021.",
]:
    add_list_item(doc, item, 20)

add_heading(doc, "3. Work Completed")
for item in [
    "Uploaded and opened test attachments successfully while Manage was using the NFS /doclinks mount.",
    "Confirmed the PowerScale S3 endpoint, bucket, access credentials, and certificate chain required for the test.",
    "Ran IBM's file2s3.sh/FileToS3Move tool with a DOCINFOID filter so only the controlled records were migrated.",
    "Confirmed the tool reported completion without errors and updated each selected DOCINFO record.",
    "Confirmed the selected URLNAME values changed from /doclinks/... to cos:doclinks/... in Oracle.",
    "Configured the seven noncredential S3 properties and entered the access key and secret key through the Manage System Properties application.",
    "Kept credentials out of this document and Jira evidence.",
]:
    add_list_item(doc, item, 21)

add_heading(doc, "4. S3 Configuration")
properties = doc.add_table(rows=1, cols=3)
properties.style = "Table Grid"
headers = properties.rows[0].cells
for index, value in enumerate(("Property", "Expected value", "Handling")):
    set_cell_shading(headers[index], TEAL)
    set_cell_text(headers[index], value, bold=True, color=WHITE, size=8.5)
set_repeat_table_header(properties.rows[0])
rows = [
    ("mxe.cosaccesskey", "Configured", "Entered in Manage UI; value not recorded"),
    ("mxe.cossecretkey", "Configured", "Entered in Manage UI; value not recorded"),
    ("mxe.cosendpointuri", "https://bhm-pwrsclnfs.lac1.biz:9021", "Noncredential"),
    ("mxe.cosbucketname", "dr-maximo-bckt", "Noncredential"),
    ("mxe.attachmentstorage", "com.ibm.tivoli.maximo.oslc.provider.COSAttachmentStorage", "S3 provider"),
    ("mxe.doclink.securedAttachment", "true", "Required for S3"),
    ("mxe.doclink.doctypes.defpath", "cos:doclinks/default", "IBM fixed S3 path"),
    ("mxe.doclink.doctypes.topLevelPaths", "cos:doclinks", "IBM fixed S3 path"),
    ("mxe.doclink.path01", "cos:doclinks=https://drgitopswks.manage.drgitopsapp.apps.drroc4.lac1.biz/maximo/oslc/cosdoclink", "Manage UI route"),
]
for prop, expected, handling in rows:
    cells = properties.add_row().cells
    set_cell_text(cells[0], prop, size=8.2)
    set_cell_text(cells[1], expected, size=8.2)
    set_cell_text(cells[2], handling, size=8.2)
set_table_geometry(properties, [2800, 4160, 2400])

note = doc.add_paragraph()
note.paragraph_format.space_before = Pt(4)
lead = note.add_run("Note: ")
lead.bold = True
lead.font.color.rgb = RGBColor.from_string(GOLD)
note.add_run(
    "mxe.cosnestedfile is not required for this sample because the migrated object keys were flattened to the S3 root. Enable it only if migrated objects retain subfolder paths."
)

add_heading(doc, "5. Post-Restart Validation")
validation = doc.add_table(rows=1, cols=3)
validation.style = "Table Grid"
for index, value in enumerate(("Validation", "Status", "Evidence")):
    set_cell_shading(validation.rows[0].cells[index], NAVY)
    set_cell_text(validation.rows[0].cells[index], value, bold=True, color=WHITE, size=8.5)
set_repeat_table_header(validation.rows[0])
checks = [
    ("All nine properties show the expected Current Value after all Manage bundles restart", "Pending", "System Properties screenshot"),
    ("Open migrated PDF, image, text, CSV, and XML attachments", "Pending", "Manage UI test"),
    ("Test at least one Work Order, Asset, and Location", "Pending", "Manage UI test"),
    ("Upload, download, and delete one new S3 attachment", "Pending", "Manage UI and S3 listing"),
    ("New DOCINFO.URLNAME starts with cos:doclinks", "Pending", "Oracle query"),
    ("New object is visible in dr-maximo-bckt", "Pending", "PowerScale/s3cmd listing"),
    ("Manage logs show no TLS, authentication, bucket, or object errors", "Pending", "Bundle logs"),
]
for check, status, evidence in checks:
    cells = validation.add_row().cells
    set_cell_text(cells[0], check, size=8.5)
    set_cell_shading(cells[1], PALE_GOLD)
    set_cell_text(cells[1], status, bold=True, color=GOLD, size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(cells[2], evidence, size=8.5)
set_table_geometry(validation, [5000, 1300, 3060])

add_heading(doc, "6. Rollback Test")
doc.add_paragraph(
    "Rollback must be tested before this work is considered complete. Keep the original NFS files and the S3 objects during the test; rollback should change references, not delete evidence."
)

add_heading(doc, "Preparation", level=2)
for item in [
    "Export the nine current S3 property values and the previous NFS property values.",
    "Keep a mapping of each test DOCINFOID, original /doclinks URLNAME, and migrated cos:doclinks URLNAME.",
    "Confirm the original files still exist on NFS and that the /doclinks PVC can be mounted by every Manage bundle.",
    "Pause new attachment uploads during the rollback window.",
]:
    add_list_item(doc, item, 20)

add_heading(doc, "Rollback Procedure", level=2)
rollback_steps = [
    "Select three to five migrated records across Work Orders, Assets, and Locations.",
    "In Manage System Properties, restore the exact pre-test NFS values. Do not guess the previous attachment provider or path values.",
    "Restore the selected DOCINFO.URLNAME values from the saved pre-migration mapping so they point to their original /doclinks paths.",
    "Restart all Manage server bundle pods: UI, cron, MEA, report, and any additional configured bundle.",
    "Open and download the selected attachments from Manage, then upload one new NFS attachment and confirm its URLNAME uses /doclinks.",
    "Reapply the nine S3 properties through the Manage UI, restore the selected cos:doclinks URLNAME values, restart the bundles, and verify S3 access again.",
]
for item in rollback_steps:
    add_list_item(doc, item, 22)

rollback_callout = doc.add_table(rows=1, cols=1)
rollback_callout.style = "Table Grid"
set_table_geometry(rollback_callout, [9360])
set_cell_shading(rollback_callout.cell(0, 0), PALE_GOLD)
set_cell_text(
    rollback_callout.cell(0, 0),
    "Important: A full rollback after cutover is not only a property change. Any attachment created only in S3 after cutover must first be copied back to NFS, and its DOCINFO.URLNAME must be restored to the correct NFS path.",
    bold=True,
    color=NAVY,
    size=9,
)

add_heading(doc, "Rollback Success Criteria", level=2)
for item in [
    "Selected attachments open and download from NFS after rollback.",
    "A new attachment can be created on NFS.",
    "No DOCINFO record points to a file that does not exist.",
    "The environment can be returned to S3 and the same attachments remain available.",
]:
    add_list_item(doc, item, 20)

add_heading(doc, "7. Jira Closure Criteria")
for item in [
    "All post-restart validation checks pass.",
    "Rollback and return-to-S3 tests pass.",
    "Screenshots or query results are attached to Jira as evidence.",
    "No access key, secret key, Vault token, or certificate private key is attached to Jira.",
]:
    add_list_item(doc, item, 20)

add_heading(doc, "8. References")
references = [
    ("IBM - Converting file-based storage to S3", "https://www.ibm.com/docs/en/masv-and-l/maximo-manage/cd?topic=storage-converting-file-based-s3"),
    ("IBM - Configuring system properties for S3", "https://www.ibm.com/docs/en/masv-and-l/maximo-manage/cd?topic=storage-configuring-system-properties-s3"),
    ("IBM - Attachment properties for S3", "https://www.ibm.com/docs/en/masv-and-l/maximo-manage/cd?topic=properties-attachment-s3"),
    ("IBM Support - S3 storage for attachments", "https://www.ibm.com/support/pages/collecting-data-s3-storage-attachments"),
]
for label, url in references:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)
    add_hyperlink(paragraph, label, url)

doc.save(OUTPUT)
print(OUTPUT)
