from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = "PowerScale_NFS_to_S3_MAS_Manage_Attachments_Research.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(9)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


def add_section(doc, title):
    doc.add_heading(title, level=1)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(10)
styles["Normal"].paragraph_format.space_after = Pt(6)

for style_name, size, color in [
    ("Heading 1", 14, RGBColor(31, 78, 121)),
    ("Heading 2", 12, RGBColor(31, 78, 121)),
]:
    style = styles[style_name]
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = color
    style.paragraph_format.space_before = Pt(10)
    style.paragraph_format.space_after = Pt(4)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("PowerScale NFS to PowerScale S3 Research\nfor MAS Manage Attachments")
run.bold = True
run.font.name = "Arial"
run.font.size = Pt(18)
run.font.color.rgb = RGBColor(31, 78, 121)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.add_run("Research document for Jira attachment").italic = True

add_section(doc, "Summary")
doc.add_paragraph(
    "This document captures initial research for moving IBM MAS Manage attachments "
    "from the current PowerScale NFS/file-based storage model to PowerScale "
    "S3-compatible object storage. This is research only; no implementation or "
    "migration is included in this ticket."
)

add_section(doc, "Current and Target State")
table = doc.add_table(rows=1, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
hdr = table.rows[0].cells
set_cell_text(hdr[0], "Current State", True)
set_cell_text(hdr[1], "Target State", True)
for cell in hdr:
    set_cell_shading(cell, "D9EAF7")
row = table.add_row().cells
set_cell_text(row[0], "MAS Manage attachments are currently stored using file-based/NFS storage on PowerScale.")
set_cell_text(row[1], "MAS Manage attachments use PowerScale S3-compatible object storage.")

add_section(doc, "IBM Supported Approach")
doc.add_paragraph(
    "IBM supports S3 storage for Maximo Manage attachments by configuring Manage "
    "system properties. The key property values and requirements are listed below."
)
add_bullets(
    doc,
    [
        "mxe.cosaccesskey",
        "mxe.cossecretkey",
        "mxe.cosendpointuri",
        "mxe.cosbucketname",
        "mxe.attachmentstorage = com.ibm.tivoli.maximo.oslc.provider.COSAttachmentStorage",
        "mxe.doclink.securedAttachment = true",
        "mxe.doclink.doctypes.defpath",
        "mxe.doclink.doctypes.topLevelPaths",
        "mxe.doclink.path01",
    ],
)

add_section(doc, "IBM Migration Tool")
doc.add_paragraph(
    "IBM provides a migration script named file2s3.sh under tools/maximo. The tool "
    "converts existing file-based attachments to S3 attachments."
)
doc.add_paragraph(
    "Important requirement: the Manage maxinst or admin pod must be able to access "
    "the existing NFS attachment folder paths."
)
params = doc.add_table(rows=1, cols=2)
params.alignment = WD_TABLE_ALIGNMENT.CENTER
params.style = "Table Grid"
hdr = params.rows[0].cells
set_cell_text(hdr[0], "Parameter", True)
set_cell_text(hdr[1], "Purpose", True)
for cell in hdr:
    set_cell_shading(cell, "D9EAF7")
for parameter, purpose in [
    ("-b", "S3 bucket name"),
    ("-x", "Access key"),
    ("-s", "Secret key"),
    ("-r", "S3 endpoint"),
    ("-w", "docinfo where clause filter"),
]:
    row = params.add_row().cells
    set_cell_text(row[0], parameter)
    set_cell_text(row[1], purpose)
doc.add_paragraph(
    'Example test command: file2s3.sh -b<bucket> -x<access-key> -s<secret-key> '
    '-r<endpoint> -w"(DOCUMENT=\'TESTDOC\')"'
)

add_section(doc, "PowerScale S3 Notes")
add_bullets(
    doc,
    [
        "PowerScale OneFS supports S3-compatible object storage.",
        "S3 buckets map to directories in OneFS.",
        "S3 object keys map to file system paths.",
        "S3 access requires an access ID and secret key.",
        "Bucket names must follow DNS-style naming rules.",
        "Bucket path and access zone configuration must be confirmed with the storage team.",
        "HTTPS endpoint and CA trust must be confirmed before testing.",
    ],
)

add_section(doc, "Items to Confirm")
add_bullets(
    doc,
    [
        "PowerScale S3 endpoint URL",
        "Bucket name for Manage attachments",
        "Access key and secret key",
        "TLS/CA certificate requirements",
        "Whether path-style S3 access is required",
        "Existing NFS attachment path",
        "Whether the Manage admin/maxinst pod can access the NFS path",
        "Whether migrated documents are stored flat or nested",
        "Whether mxe.cosnestedfile=1 is required",
        "Backup and rollback plan",
    ],
)

add_section(doc, "Risks")
add_bullets(
    doc,
    [
        "Existing NFS folder structure may not match S3 expectations.",
        "New S3 uploads are expected at the bucket root.",
        "Migrated documents in subfolders may require mxe.cosnestedfile=1.",
        "Any integration that directly reads attachments from NFS may need changes.",
        "S3 credentials must be stored securely.",
        "Migration must be tested first with a small docinfo filter.",
    ],
)

add_section(doc, "Recommended Next Step")
add_numbered(
    doc,
    [
        "Configure a non-production PowerScale S3 bucket.",
        "Configure Manage S3 attachment properties.",
        "Test new attachment upload/download.",
        "Run file2s3.sh with a small docinfo filter.",
        "Validate migrated attachments in the Manage UI.",
        "Document rollback to NFS.",
    ],
)

add_section(doc, "Closure")
doc.add_paragraph("This research ticket can be closed after this document is attached to the Jira.")

add_section(doc, "References")
add_bullets(
    doc,
    [
        "IBM: Converting file-based storage to S3 storage - https://www.ibm.com/docs/en/masv-and-l/maximo-manage/cd?topic=storage-converting-file-based-s3",
        "IBM: Attachment properties for S3 storage - https://www.ibm.com/docs/en/masv-and-l/maximo-manage/cd?topic=properties-attachment-s3",
        "IBM Support: Collecting Data - S3 Storage for Attachments - https://www.ibm.com/support/pages/collecting-data-s3-storage-attachments",
        "Dell: PowerScale OneFS S3 API Reference - https://www.dell.com/support/manuals/en-us/isilon-onefs/ifs_pub_onefs_api_reference/s3-buckets-resource",
        "Dell: PowerScale OneFS S3 Overview - https://infohub.delltechnologies.com/en-ca/l/dell-powerscale-onefs-s3-overview/buckets-8/",
    ],
)

doc.save(OUT)
print(OUT)
